"""Bulk learner import / export from Excel (apps/admissions/bulk.py).

The office fills in the template and imports it; the platform creates the
learner accounts, applications, parents and enrolments; an export comes back
in the same format and re-imports without duplicating anyone.
"""
import io
import tempfile
from datetime import date
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import ParentLink, Person, PersonContact
from apps.learning.models import ModuleEnrolment, ProgrammeEnrolment
from core import academic_spine, school

from . import bulk
from .models import Application, Guardian

User = get_user_model()


def sheet(*rows):
    """xlsx bytes of the import format holding ``rows`` (``{column key: value}``)."""
    buffer = io.BytesIO()
    bulk.build_workbook(rows, examples=False).save(buffer)
    return buffer.getvalue()


def young_learner(**extra):
    row = {
        'first_name': 'Thandi', 'last_name': 'Mokoena', 'gender': 'Female',
        'date_of_birth': date(2019, 4, 12), 'grade': 1, 'city': 'Johannesburg',
        'suburb': 'Yeoville', 'province': 'Gauteng', 'address': '12 Raleigh Street',
        'id_document_type': 'Birth certificate', 'id_number': '1904120800085',
        'home_language': 'isiZulu', 'writing_hand': 'Right', 'paid_up_to': f'{school.YEAR}-03',
        'g1_role': 'Mother', 'g1_title': 'Mrs', 'g1_name': 'Nomsa Mokoena',
        'g1_cell': '082 555 0101', 'g1_email': 'nomsa@example.com',
        'g2_role': 'Father', 'g2_name': 'Sipho Mokoena', 'g2_email': 'sipho@example.com',
        'emergency_name': 'Grace Dlamini', 'has_medical_condition': 'Yes',
        'medical_conditions': 'Asthma', 'medical_aid_name': 'Discovery', 'fee_payer': 'Both parents',
        'declarations': 'Yes', 'office_account_number': 'MOK001',
    }
    row.update(extra)
    return row


def fet_learner(**extra):
    row = {
        'first_name': 'Kwame', 'last_name': 'Mensah', 'gender': 'Male',
        'date_of_birth': date(2010, 1, 27), 'email': 'kwame@example.com', 'grade': 10,
        'class_code': '10A', 'is_new_learner': 'Yes', 'fal': 'Afrikaans First Additional Language',
        'maths': 'Mathematical Literacy', 'elective1': 'Physical Sciences',
        'elective2': 'Life Sciences', 'elective3': 'Accounting', 'is_south_african': 'No',
        'id_document_type': 'Passport', 'passport_number': 'G1234567',
        'document_expiry': date(2030, 5, 31),
        'g1_role': 'Legal guardian', 'g1_name': 'Kofi Mensah', 'g1_email': 'kofi@example.com',
    }
    row.update(extra)
    return row


class BulkTestCase(TestCase):

    @classmethod
    def setUpTestData(cls):
        academic_spine.seed(verbose=False, calendar=False)

    def run_import(self, *rows, **options):
        options.setdefault('dry_run', False)
        return bulk.import_workbook(sheet(*rows), **options)


class TemplateTests(BulkTestCase):

    def test_template_generates_with_examples_instructions_and_subjects(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'template.xlsx'
            call_command('learners_template', out=str(path), stdout=io.StringIO())
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames[:3], ['Learners', 'Instructions', 'Subjects'])
            ws = wb['Learners']
            headers = [ws.cell(row=bulk.HEADER_ROW, column=i).value
                       for i in range(1, len(bulk.COLUMNS) + 1)]
            self.assertEqual(headers, [c.header for c in bulk.COLUMNS])
            self.assertEqual(ws.freeze_panes, 'C3')
            self.assertTrue(ws.data_validations.dataValidation)
            firsts = [ws.cell(row=r, column=1).value for r in range(3, 6)]
            self.assertTrue(all(v.startswith('EXAMPLE') for v in firsts))
            instructions = [r[0] for r in wb['Instructions'].iter_rows(values_only=True)]
            self.assertIn('Grade', instructions)
            subjects = list(wb['Subjects'].iter_rows(min_row=2, values_only=True))
            self.assertIn((10, 'Physical Sciences', 'PHYS-SCI', 'Elective subjects — choose three'),
                          subjects)

            # The examples are valid and skipped on import; with the prefix
            # removed they import cleanly.
            report = bulk.import_workbook(str(path))
            self.assertEqual(report.counts['skipped'], 3)
            self.assertEqual(report.counts['created'], 0)
            for r in range(3, 6):
                cell = ws.cell(row=r, column=1)
                cell.value = cell.value.replace('EXAMPLE ', '')
            buffer = io.BytesIO()
            wb.save(buffer)
            report = bulk.import_workbook(buffer.getvalue(), dry_run=True)
            self.assertEqual(report.counts, {'created': 3, 'updated': 0, 'skipped': 0, 'error': 0},
                             report.as_dict())
        self.assertFalse(Person.objects.filter(first_name='Thandi').exists())  # dry run

    def test_static_template_is_present(self):
        from django.conf import settings
        path = Path(settings.BASE_DIR) / 'static' / 'documents' / 'UCS-learner-import-template.xlsx'
        self.assertTrue(path.exists())


class ImportTests(BulkTestCase):

    def test_import_creates_learners_guardians_and_enrolments(self):
        report = self.run_import(young_learner(), fet_learner())
        self.assertTrue(report.ok, report.as_dict())
        self.assertEqual(report.counts['created'], 2)

        # Young learner — no e-mail: a generated, non-routable login.
        thandi = Person.objects.get(first_name='Thandi')
        self.assertEqual(thandi.user.email, 'thandi.mokoena.1@learners.ucs.local')
        self.assertIn('generated', ' '.join(m for _, m in report.rows[0].warnings))
        self.assertEqual(thandi.user_type, 'student')
        self.assertTrue(thandi.registered and thandi.profile_status)
        self.assertEqual(thandi.date_of_birth, date(2019, 4, 12))
        self.assertEqual(PersonContact.objects.get(person=thandi).city, 'Johannesburg')
        app = Application.objects.get(person=thandi, year=school.YEAR)
        self.assertEqual(app.status, Application.STATUS_ADMITTED)
        self.assertEqual(app.programme.grade, 1)
        self.assertEqual(app.id_document_type, Application.DOC_BIRTH_CERT)
        self.assertTrue(app.has_medical_condition)
        self.assertTrue(app.accept_terms and app.accept_popia)
        self.assertEqual(app.office_account_number, 'MOK001')
        self.assertEqual(app.fee_payer, 'both')
        self.assertEqual(set(app.guardians.values_list('role', flat=True)),
                         {Guardian.ROLE_MOTHER, Guardian.ROLE_FATHER})
        enrolment = ProgrammeEnrolment.objects.get(person=thandi)
        self.assertEqual(enrolment.cohort.code, str(school.YEAR))
        modules = ModuleEnrolment.objects.filter(person=thandi)
        self.assertEqual(modules.count(), 5)
        # "Fees paid up to" March → unlocked to the end of March.
        self.assertTrue(all(m.status == ModuleEnrolment.STATUS_ACTIVE
                            and m.paid_until == date(school.YEAR, 3, 31) for m in modules))

        # FET learner — chosen subjects, class 10A, LOCKED (no fees column).
        kwame = Person.objects.get(user__email='kwame@example.com')
        self.assertEqual(ProgrammeEnrolment.objects.get(person=kwame).cohort.code, '10A')
        codes = set(ModuleEnrolment.objects.filter(person=kwame)
                    .values_list('programme_module__code', flat=True))
        self.assertEqual(codes, {'ENG-HL', 'LO', 'AFR-FAL', 'MLIT', 'PHYS-SCI', 'LIFE-SCI', 'ACC'})
        self.assertFalse(ModuleEnrolment.objects.filter(person=kwame).exclude(
            status=ModuleEnrolment.STATUS_LOCKED).exists())
        self.assertFalse(Application.objects.get(person=kwame).is_south_african)
        # No parent accounts unless asked for.
        self.assertFalse(ParentLink.objects.exists())

    def test_reimport_updates_instead_of_duplicating(self):
        self.run_import(young_learner(), fet_learner())
        users, apps = User.objects.count(), Application.objects.count()
        report = self.run_import(young_learner(city='Pretoria'),
                                 fet_learner(elective3='Economics', city='Soweto'))
        self.assertEqual(report.counts['updated'], 2, report.as_dict())
        self.assertEqual((User.objects.count(), Application.objects.count()), (users, apps))
        thandi = Person.objects.get(first_name='Thandi')         # matched on ID number
        self.assertEqual(thandi.contact.city, 'Pretoria')
        self.assertEqual(Guardian.objects.filter(application__person=thandi).count(), 2)
        kwame = Person.objects.get(user__email='kwame@example.com')
        codes = set(ModuleEnrolment.objects.filter(person=kwame)
                    .values_list('programme_module__code', flat=True))
        self.assertIn('ECON', codes)
        self.assertNotIn('ACC', codes)                           # dropped choice removed

    def test_invalid_grade_and_subject_choice_reported_with_row_and_column(self):
        report = self.run_import(
            young_learner(grade=13),
            fet_learner(elective1='Mathematics', email='a@example.com'),
            fet_learner(elective2='Physical Sciences', email='b@example.com'),
            fet_learner(fal='', email='c@example.com'),
            young_learner(first_name='Ayanda', id_number='1501010800089', gender='Purple',
                          date_of_birth='31/02/2015'),
        )
        rows = {r.row: r for r in report.rows}
        self.assertEqual(report.counts['error'], 5)

        def columns(row):
            return {column for column, _ in rows[row].errors}
        self.assertIn('Grade', columns(3))
        self.assertIn('Elective 1 (Gr 10-12)', columns(4))
        self.assertIn('not offered in Grade 10', ' '.join(m for _, m in rows[4].errors))
        self.assertIn('Elective 3 (Gr 10-12)', columns(5))           # duplicate electives
        self.assertIn('FAL (Gr 10-12)', columns(6))                  # incomplete choices
        self.assertEqual(columns(7), {'Gender', 'Date of birth'})
        self.assertFalse(Person.objects.filter(user_type='student').exists())

    def test_one_bad_row_does_not_stop_the_rest(self):
        report = self.run_import(young_learner(grade=0), fet_learner())
        self.assertEqual(report.counts, {'created': 1, 'updated': 0, 'skipped': 0, 'error': 1})
        self.assertTrue(Person.objects.filter(user__email='kwame@example.com').exists())

    def test_dry_run_saves_nothing(self):
        report = self.run_import(young_learner(), fet_learner(), dry_run=True,
                                 create_parent_accounts=True)
        self.assertEqual(report.counts['created'], 2)
        self.assertFalse(Application.objects.exists())
        self.assertFalse(User.objects.filter(email__in=['kwame@example.com',
                                                        'nomsa@example.com']).exists())

    def test_fet_without_choices_gets_default_subjects_and_a_warning(self):
        report = self.run_import(fet_learner(fal='', maths='', elective1='', elective2='',
                                             elective3=''))
        self.assertTrue(report.ok, report.as_dict())
        self.assertIn('default subjects', ' '.join(m for _, m in report.rows[0].warnings))
        kwame = Person.objects.get(user__email='kwame@example.com')
        self.assertEqual(ModuleEnrolment.objects.filter(person=kwame).count(), 7)

    def test_parent_shared_by_two_children_links_once(self):
        sister = young_learner(first_name='Lerato', date_of_birth=date(2015, 8, 3), grade=5,
                               id_number='1508030800082', g2_name='', g2_role='', g2_email='')
        report = self.run_import(young_learner(), sister, create_parent_accounts=True)
        self.assertTrue(report.ok, report.as_dict())
        nomsa = User.objects.get(email='nomsa@example.com')
        self.assertEqual(User.objects.filter(email='nomsa@example.com').count(), 1)
        self.assertEqual(nomsa.profile.user_type, 'parent')
        self.assertEqual(ParentLink.objects.filter(parent=nomsa).count(), 2)
        self.assertEqual(ParentLink.objects.filter(parent__email='sipho@example.com').count(), 1)
        # Re-importing does not link again.
        self.run_import(young_learner(), sister, create_parent_accounts=True)
        self.assertEqual(ParentLink.objects.filter(parent=nomsa).count(), 2)

    def test_parent_invites_in_dry_run_are_only_reported(self):
        report = self.run_import(young_learner(), dry_run=True, send_invites=True)
        self.assertIn('would be sent', ' '.join(report.rows[0].notes))

    def test_not_a_workbook(self):
        report = bulk.import_workbook(b'not an excel file')
        self.assertTrue(report.fatal)


class ExportTests(BulkTestCase):

    def test_export_reimport_round_trip(self):
        self.run_import(young_learner(), fet_learner(),
                        young_learner(first_name='Lerato', date_of_birth=date(2015, 8, 3), grade=5,
                                      id_number='', paid_up_to=''))
        first = bulk.export_rows(year=school.YEAR)
        self.assertEqual(len(first), 3)
        data = bulk.export_workbook(year=school.YEAR)
        counts = (User.objects.count(), Application.objects.count(), Guardian.objects.count(),
                  ModuleEnrolment.objects.count())
        report = bulk.import_workbook(data, dry_run=False)
        self.assertEqual(report.counts, {'created': 0, 'updated': 3, 'skipped': 0, 'error': 0},
                         report.as_dict())
        self.assertEqual((User.objects.count(), Application.objects.count(),
                          Guardian.objects.count(), ModuleEnrolment.objects.count()), counts)
        self.assertEqual(bulk.export_rows(year=school.YEAR), first)

        # The export carries the same headers as the template, and the data.
        ws = load_workbook(io.BytesIO(data))['Learners']
        self.assertEqual(ws.cell(row=bulk.HEADER_ROW, column=1).value, 'First name(s)')
        by_name = {r['first_name']: r for r in first}
        self.assertEqual(by_name['Thandi']['email'], 'thandi.mokoena.1@learners.ucs.local')
        self.assertEqual(by_name['Thandi']['paid_up_to'], f'{school.YEAR}-03')
        self.assertEqual(by_name['Kwame']['maths'], 'Mathematical Literacy')
        self.assertEqual(by_name['Kwame']['class_code'], '10A')
        self.assertEqual(by_name['Thandi']['g1_name'], 'Nomsa Mokoena')

    def test_grade_filter_and_command(self):
        self.run_import(young_learner(), fet_learner())
        self.assertEqual([r['first_name'] for r in bulk.export_rows(year=school.YEAR, grade=10)],
                         ['Kwame'])
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / 'g10.xlsx'
            call_command('learners_export', year=school.YEAR, grade=10, out=str(out),
                         stdout=io.StringIO())
            ws = load_workbook(out)['Learners']
            self.assertEqual(ws.cell(row=bulk.FIRST_DATA_ROW, column=1).value, 'Kwame')

    def test_enrolled_learner_without_application_is_exported(self):
        user = User.objects.create_user(username='demo@example.com', email='demo@example.com')
        person = user.profile
        person.first_name, person.last_name, person.gender = 'Demo', 'Learner', 'male'
        person.date_of_birth = date(2014, 2, 2)
        person.save()
        from apps.learning.models import Programme
        academic_spine.enrol_student(person, Programme.objects.get(institution__code='UCS', grade=6),
                                     activate=False)
        rows = bulk.export_rows(year=school.YEAR)
        self.assertEqual([(r['first_name'], r['grade']) for r in rows], [('Demo', 6)])


class PageTests(BulkTestCase):

    def setUp(self):
        self.staff = User.objects.create_user(username='office@example.com',
                                              email='office@example.com', password='x-Pass-1234',
                                              is_staff=True)
        Person.objects.filter(user=self.staff).update(user_type='staff', registered=True,
                                                      profile_status=True)
        self.learner = User.objects.create_user(username='l@example.com', email='l@example.com',
                                                password='x-Pass-1234')
        Person.objects.filter(user=self.learner).update(user_type='student', registered=True,
                                                        profile_status=True)

    def test_non_staff_gets_404(self):
        self.client.force_login(self.learner)
        for name in ('admissions:bulk-import', 'admissions:bulk-export', 'admissions:bulk-template'):
            self.assertEqual(self.client.get(reverse(name)).status_code, 404, name)
        upload = SimpleUploadedFile('l.xlsx', sheet(fet_learner()))
        self.assertEqual(self.client.post(reverse('admissions:bulk-import'),
                                          {'file': upload}).status_code, 404)
        self.assertFalse(Application.objects.exists())

    def test_staff_preview_then_confirm(self):
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse('admissions:bulk-import')).status_code, 200)
        template = self.client.get(reverse('admissions:bulk-template'))
        self.assertEqual(template.status_code, 200)
        self.assertIn('spreadsheetml', template['Content-Type'])

        upload = SimpleUploadedFile('learners.xlsx', sheet(young_learner(), fet_learner(grade=99)))
        response = self.client.post(reverse('admissions:bulk-import'),
                                    {'action': 'preview', 'file': upload})
        self.assertContains(response, 'Preview — dry run')
        self.assertContains(response, 'will be created')
        self.assertContains(response, 'error — not imported')
        self.assertFalse(Application.objects.exists())

        response = self.client.post(reverse('admissions:bulk-import'),
                                    {'action': 'import', 'create_parent_accounts': '1'})
        self.assertContains(response, 'Import complete')
        self.assertEqual(Application.objects.count(), 1)
        self.assertEqual(ParentLink.objects.count(), 2)
        app = Application.objects.get()
        self.assertEqual(app.decided_by, self.staff)

    def test_staff_export_download(self):
        self.client.force_login(self.staff)
        bulk.import_workbook(sheet(fet_learner()), dry_run=False)
        self.assertEqual(self.client.get(reverse('admissions:bulk-export')).status_code, 200)
        response = self.client.get(reverse('admissions:bulk-export'),
                                   {'year': school.YEAR, 'grade': 10, 'download': '1'})
        self.assertEqual(response.status_code, 200)
        ws = load_workbook(io.BytesIO(response.content))['Learners']
        self.assertEqual(ws.cell(row=bulk.FIRST_DATA_ROW, column=1).value, 'Kwame')

    def test_office_list_links_to_import(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse('admissions:office-list'))
        self.assertContains(response, reverse('admissions:bulk-import'))
