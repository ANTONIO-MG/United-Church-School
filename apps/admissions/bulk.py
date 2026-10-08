"""Bulk import and export of learners from an Excel workbook.

The school office keeps its learner register in a spreadsheet. This module
moves that register in and out of the platform in ONE column format:

* :data:`COLUMNS` — the specification of every column, in sheet order: its
  header, the section it belongs to, the help text the Instructions sheet
  shows, whether it is required, the values it accepts and the model field it
  maps to. Everything else here (the template, the import, the export) is
  driven from that list, so a column is added in one place.
* :func:`build_workbook` / :func:`write_template` — the import template:
  a styled "Learners" sheet with drop-downs, three EXAMPLE rows, an
  "Instructions" sheet explaining every column and a "Subjects" sheet listing
  the subjects per grade. ``python manage.py learners_template`` writes it to
  ``static/documents/UCS-learner-import-template.xlsx``.
* :func:`import_workbook` — reads a filled-in sheet and creates or updates,
  per row: the learner's account (User + Person + PersonContact), the
  Application for the school year with its Guardians, and the enrolment in the
  grade and class with its subjects. Returns an :class:`ImportReport`.
* :func:`export_workbook` — writes learners from the database in exactly the
  same format, so an export can be edited and imported again.

Import rules
------------
* **Matching** (idempotent): a row updates the learner with the same login
  e-mail; with no e-mail, the learner whose application carries the same ID /
  birth-certificate number; failing that, the student with the same first
  name, surname and date of birth. Otherwise a new learner is created.
* **Blank cells keep** what is already recorded for an existing learner; for a
  new learner they take the default the Instructions sheet names.
* **No e-mail** (young learners): a unique, non-routable login
  ``firstname.surname.N@learners.ucs.local`` is generated and the report says
  so. The export writes it back, so a re-import matches on it.
* **Subjects** are left LOCKED (fees unpaid) unless *Fees paid up to* names a
  month: then every subject in the grade is unlocked to the end of that month.
* Rows whose first name starts with ``EXAMPLE`` are skipped. A dry run
  (the default) does everything inside a transaction that is rolled back, so
  the preview is exactly what the real import will do. A real import commits
  each row on its own: one bad row never stops the rest.
"""
import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime

from django.contrib.auth import get_user_model
from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from core import school

from .models import REQUIRED_DECLARATIONS, Application, Guardian

logger = logging.getLogger('admissions')

#: Domain of the generated logins of learners who have no e-mail address.
GENERATED_EMAIL_DOMAIN = 'learners.ucs.local'
#: Rows whose first name starts with this are examples and never imported.
EXAMPLE_PREFIX = 'EXAMPLE'
SHEET_LEARNERS = 'Learners'
SHEET_INSTRUCTIONS = 'Instructions'
SHEET_SUBJECTS = 'Subjects'
SHEET_LISTS = 'Lists'
#: Template rows given drop-downs and text formatting.
TEMPLATE_ROWS = 1000

NAVY = '00498B'
GOLD = 'F2C14E'

# ---------------------------------------------------------------------------
# Value lists
# ---------------------------------------------------------------------------
YES_NO = ((True, 'Yes'), (False, 'No'))
GENDERS = (('male', 'Male'), ('female', 'Female'))
PROVINCES = tuple((p, p) for p in (
    'Gauteng', 'Eastern Cape', 'Free State', 'KwaZulu-Natal', 'Limpopo', 'Mpumalanga',
    'North West', 'Northern Cape', 'Western Cape', 'Outside South Africa'))
HOME_LANGUAGES = tuple((x, x) for x in (
    'English', 'isiZulu', 'isiXhosa', 'Afrikaans', 'Sesotho', 'Setswana', 'Sepedi', 'Xitsonga',
    'siSwati', 'Tshivenda', 'isiNdebele', 'French', 'Portuguese', 'Shona', 'Swahili', 'Other'))
RELIGIONS = tuple((x, x) for x in (
    'Christian', 'Muslim', 'Hindu', 'Jewish', 'African traditional', 'None', 'Other'))
STATUSES = tuple((value, label.split(' — ')[0]) for value, label in Application.STATUS_CHOICES)
ID_DOCUMENTS = (
    (Application.DOC_BIRTH_CERT, 'Birth certificate'),
    (Application.DOC_SA_ID, 'SA ID'),
    (Application.DOC_PASSPORT, 'Passport'),
    (Application.DOC_ASYLUM, 'Asylum / refugee permit'),
)
RACES = tuple((v, label) for v, label in Application.RACE_CHOICES)
HANDS = tuple(Application.HAND_CHOICES)
FEE_PAYERS = tuple(Application.FEE_PAYER_CHOICES)
ROLES = tuple(Guardian.ROLE_CHOICES)
TITLES = tuple((v, label) for v, label in Guardian.TITLE_CHOICES if v)
GRADES = tuple((g, str(g)) for g in school.GRADES)

#: Aliases accepted for some choice values, on top of the value and its label.
ALIASES = {
    'yes': True, 'y': True, 'true': True, '1': True, 'ja': True,
    'no': False, 'n': False, 'false': False, '0': False, 'nee': False,
}


def _subject_names(group):
    names = {s['code']: s['name'] for s in school.SUBJECTS}
    return tuple((code, names[code]) for code, g, _ in school.subjects_for(10) if g == group)


FAL_SUBJECTS = _subject_names(school.GROUP_FAL)
MATHS_SUBJECTS = _subject_names(school.GROUP_MATHS)
ELECTIVE_SUBJECTS = _subject_names(school.GROUP_ELECTIVE)


# ---------------------------------------------------------------------------
# The column specification
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Column:
    """One column of the learner sheet.

    ``target`` is ``<object>.<field>`` — ``user``, ``person``, ``contact``,
    ``app`` (the Application), ``g1`` / ``g2`` (the first / second parent) — or
    a special key handled by the importer (``grade``, ``cohort``, ``subject``,
    ``paid_up_to``, ``declarations``). ``kind`` drives parsing and the
    drop-down: text, email, phone, int, date, month, bool, choice, grade,
    subject. A ``choice`` column with ``strict=False`` only *suggests* values.
    """
    key: str
    header: str
    section: str
    target: str
    kind: str = 'text'
    required: bool = False
    help: str = ''
    choices: tuple = ()
    strict: bool = True
    default: object = None
    width: int = 18
    group: str = ''           # subject columns: the choice group

    @property
    def obj(self):
        return self.target.split('.', 1)[0]

    @property
    def attr(self):
        return self.target.split('.', 1)[1] if '.' in self.target else ''

    def labels(self):
        return [label for _, label in self.choices if label]


def _parent_columns(n):
    p = f'g{n}'
    s = f'Parent / guardian {n}'
    return [
        Column(f'{p}_role', f'Parent {n} role', s, f'{p}.role', 'choice', choices=ROLES,
               help='Father, Mother or Legal guardian. Required when a parent name is given; the '
                    'two parents must have different roles.', width=14),
        Column(f'{p}_title', f'Parent {n} title', s, f'{p}.title', 'choice', choices=TITLES,
               width=10, help='Mr, Mrs, Ms, Miss, Dr, Prof, Pastor or Rev.'),
        Column(f'{p}_name', f'Parent {n} full name and surname', s, f'{p}.full_name', width=26,
               help='Leave the whole parent block blank when there is no such parent.'),
        Column(f'{p}_id', f'Parent {n} ID / passport number', s, f'{p}.id_number', width=20),
        Column(f'{p}_cell', f'Parent {n} cell number', s, f'{p}.cell_phone', 'phone', width=16),
        Column(f'{p}_home', f'Parent {n} home number', s, f'{p}.home_phone', 'phone', width=16),
        Column(f'{p}_work', f'Parent {n} work number', s, f'{p}.work_phone', 'phone', width=16),
        Column(f'{p}_email', f'Parent {n} e-mail', s, f'{p}.email', 'email', width=26,
               help='Used for the parent account / invitation. One parent with several children at '
                    'UCS uses the same e-mail on each child\'s row and gets one account.'),
        Column(f'{p}_address', f'Parent {n} residential address', s, f'{p}.residential_address',
               width=30),
        Column(f'{p}_occupation', f'Parent {n} occupation', s, f'{p}.occupation'),
        Column(f'{p}_employer', f'Parent {n} employer', s, f'{p}.employer'),
        Column(f'{p}_employer_phone', f'Parent {n} employer contact number', s,
               f'{p}.employer_phone', 'phone', width=16),
    ]


COLUMNS = [
    # -- the learner ----------------------------------------------------------
    Column('first_name', 'First name(s)', 'Learner', 'person.first_name', required=True, width=18,
           help='As on the birth certificate / ID. Rows starting with "EXAMPLE" are skipped.'),
    Column('last_name', 'Surname', 'Learner', 'person.last_name', required=True, width=18),
    Column('gender', 'Gender', 'Learner', 'person.gender', 'choice', required=True,
           choices=GENDERS, width=10),
    Column('date_of_birth', 'Date of birth', 'Learner', 'person.date_of_birth', 'date',
           required=True, width=14, help='YYYY-MM-DD, e.g. 2015-03-21 (an Excel date is fine).'),
    Column('email', 'Learner e-mail (login)', 'Learner', 'user.email', 'email', width=30,
           help='The learner\'s own login e-mail. Optional for young learners: if blank, a login '
                f'like thandi.mokoena.1@{GENERATED_EMAIL_DOMAIN} is generated (it cannot receive '
                'mail). Matching on re-import uses this e-mail.'),
    Column('phone', 'Learner cell number', 'Learner', 'person.phone', 'phone', width=16,
           help='Optional. Type numbers as text so the leading 0 is kept, e.g. 082 123 4567.'),
    Column('address', 'Physical address', 'Learner', 'app.physical_address', width=30,
           help='Street address where the learner lives.'),
    Column('suburb', 'Suburb', 'Learner', 'contact.suburb'),
    Column('city', 'City / town', 'Learner', 'contact.city'),
    Column('province', 'Province', 'Learner', 'contact.province', 'choice', choices=PROVINCES,
           strict=False, width=16),
    # -- school ---------------------------------------------------------------
    Column('year', 'School year', 'School', 'app.year', 'int', default=school.YEAR, width=10,
           help=f'The school year this row is for. Blank = {school.YEAR}.'),
    Column('grade', 'Grade', 'School', 'grade', 'grade', required=True, choices=GRADES, width=8,
           help='1 to 12.'),
    Column('class_code', 'Class', 'School', 'cohort', width=10,
           help='The class code within the grade, e.g. 10A. Blank = the year\'s class (e.g. '
                '"2026"). A class that does not exist yet is created.'),
    Column('is_new_learner', 'New to UCS', 'School', 'app.is_new_learner', 'bool',
           choices=YES_NO, default=False, width=10,
           help='Yes = joining UCS this year; No = already a UCS learner. Blank = No.'),
    Column('status', 'Application status', 'School', 'app.status', 'choice', choices=STATUSES,
           default=Application.STATUS_ADMITTED, width=16,
           help='Blank = Admitted (an existing learner). Declined / Withdrawn learners are not '
                'active in the grade.'),
    Column('fal', 'FAL (Gr 10-12)', 'School', 'subject', 'subject', choices=FAL_SUBJECTS,
           group=school.GROUP_FAL, width=24,
           help='Grade 10 – 12 only: the First Additional Language. Fill in all five subject '
                'choices, or leave them all blank for the default subjects.'),
    Column('maths', 'Maths (Gr 10-12)', 'School', 'subject', 'subject', choices=MATHS_SUBJECTS,
           group=school.GROUP_MATHS, width=22,
           help='Grade 10 – 12 only: Mathematics or Mathematical Literacy.'),
    Column('elective1', 'Elective 1 (Gr 10-12)', 'School', 'subject', 'subject',
           choices=ELECTIVE_SUBJECTS, group=school.GROUP_ELECTIVE, width=20,
           help='Grade 10 – 12 only: three different electives.'),
    Column('elective2', 'Elective 2 (Gr 10-12)', 'School', 'subject', 'subject',
           choices=ELECTIVE_SUBJECTS, group=school.GROUP_ELECTIVE, width=20),
    Column('elective3', 'Elective 3 (Gr 10-12)', 'School', 'subject', 'subject',
           choices=ELECTIVE_SUBJECTS, group=school.GROUP_ELECTIVE, width=20),
    Column('paid_up_to', 'Fees paid up to (month)', 'School', 'paid_up_to', 'month', width=14,
           help='YYYY-MM, e.g. 2026-03. The learner\'s subjects are unlocked to the end of that '
                'month. Blank = subjects stay LOCKED until fees are paid on the platform.'),
    Column('highest_grade_passed', 'Highest grade passed', 'School', 'app.highest_grade_passed',
           width=12, help='e.g. Grade 4.'),
    Column('year_grade_passed', 'Year grade passed', 'School', 'app.year_grade_passed', 'int',
           width=10),
    # -- identity -------------------------------------------------------------
    Column('id_document_type', 'Identity document', 'Identity', 'app.id_document_type', 'choice',
           choices=ID_DOCUMENTS, width=18,
           help='Birth certificate (up to Grade 9), SA ID, Passport or Asylum / refugee permit.'),
    Column('id_number', 'ID / birth certificate number', 'Identity', 'app.id_number', width=18,
           help='Used to match the learner on re-import when there is no e-mail. Type as text.'),
    Column('passport_number', 'Passport number', 'Identity', 'app.passport_number'),
    Column('permit_number', 'Asylum / refugee permit number', 'Identity', 'app.permit_number'),
    Column('document_country', 'Country of issue', 'Identity', 'app.document_country'),
    Column('document_expiry', 'Passport / permit expiry', 'Identity', 'app.document_expiry',
           'date', width=14),
    Column('is_south_african', 'South African citizen', 'Identity', 'app.is_south_african', 'bool',
           choices=YES_NO, default=True, width=10, help='Blank = Yes.'),
    Column('permanent_residency', 'Permanent residency / refugee status', 'Identity',
           'app.permanent_residency', 'bool', choices=YES_NO, width=12),
    Column('study_permit_number', 'Study permit number', 'Identity', 'app.study_permit_number',
           help='Non-South-African learners.'),
    Column('study_permit_expiry', 'Study permit expiry', 'Identity', 'app.study_permit_expiry',
           'date', width=14),
    # -- background -----------------------------------------------------------
    Column('home_language', 'Home language', 'Background', 'app.home_language', 'choice',
           choices=HOME_LANGUAGES, strict=False, width=14),
    Column('religion', 'Religion', 'Background', 'app.religion', 'choice', choices=RELIGIONS,
           strict=False, width=14),
    Column('race', 'Race (optional)', 'Background', 'app.race', 'choice', choices=RACES, width=16,
           help='Collected for Department of Education statistics only. Optional.'),
    Column('writing_hand', 'Writing hand', 'Background', 'app.writing_hand', 'choice',
           choices=HANDS, width=10),
    Column('previous_school', 'Previous school', 'Background', 'app.previous_school', width=24),
    Column('previous_school_address', 'Previous school address', 'Background',
           'app.previous_school_address', width=26),
    Column('previous_school_phone', 'Previous school contact number', 'Background',
           'app.previous_school_phone', 'phone', width=16),
    # -- family ---------------------------------------------------------------
    Column('has_father', 'Has a father', 'Family', 'app.has_father', 'bool', choices=YES_NO,
           width=10),
    Column('has_mother', 'Has a mother', 'Family', 'app.has_mother', 'bool', choices=YES_NO,
           width=10),
    Column('lives_with', 'Lives with', 'Family', 'app.lives_with', width=16,
           help='e.g. Both parents, Mother, Grandmother.'),
    *_parent_columns(1),
    *_parent_columns(2),
    # -- emergency contact ----------------------------------------------------
    Column('emergency_name', 'Emergency contact name', 'Emergency contact', 'app.emergency_name',
           width=22, help='A friend or relative other than the parents.'),
    Column('emergency_relationship', 'Emergency contact relationship', 'Emergency contact',
           'app.emergency_relationship', width=14),
    Column('emergency_home_phone', 'Emergency contact home number', 'Emergency contact',
           'app.emergency_home_phone', 'phone', width=16),
    Column('emergency_cell_phone', 'Emergency contact cell number', 'Emergency contact',
           'app.emergency_cell_phone', 'phone', width=16),
    Column('emergency_email', 'Emergency contact e-mail', 'Emergency contact',
           'app.emergency_email', 'email', width=24),
    # -- medical --------------------------------------------------------------
    Column('has_medical_condition', 'Medical condition or allergy', 'Medical',
           'app.has_medical_condition', 'bool', choices=YES_NO, width=10),
    Column('medical_conditions', 'Illnesses, allergies and medical problems', 'Medical',
           'app.medical_conditions', width=26, help='e.g. asthma, epilepsy, peanut allergy.'),
    Column('medication', 'Medication', 'Medical', 'app.medication', width=20,
           help='Medication that must be taken.'),
    Column('medical_aid_name', 'Medical aid', 'Medical', 'app.medical_aid_name'),
    Column('medical_aid_number', 'Medical aid number', 'Medical', 'app.medical_aid_number'),
    Column('medical_aid_plan', 'Medical aid plan', 'Medical', 'app.medical_aid_plan'),
    Column('medical_aid_main_member', 'Main member', 'Medical', 'app.medical_aid_main_member'),
    Column('medical_aid_main_member_phone', 'Main member contact number', 'Medical',
           'app.medical_aid_main_member_phone', 'phone', width=16),
    Column('medical_aid_main_member_id', 'Main member ID / passport number', 'Medical',
           'app.medical_aid_main_member_id'),
    Column('doctor_contact', 'Doctor (name and number)', 'Medical', 'app.doctor_contact', width=24),
    Column('medical_expenses_name', 'Responsible for medical expenses', 'Medical',
           'app.medical_expenses_name', width=22),
    Column('medical_expenses_phone', 'Medical expenses contact number', 'Medical',
           'app.medical_expenses_phone', 'phone', width=16),
    Column('medical_expenses_relationship', 'Medical expenses relationship', 'Medical',
           'app.medical_expenses_relationship', width=14),
    # -- fees -----------------------------------------------------------------
    Column('fee_payer', 'Fee payer', 'Fees', 'app.fee_payer', 'choice', choices=FEE_PAYERS,
           width=14, help='Who is responsible for paying the school fees.'),
    Column('fee_payer_name', 'Fee payer name', 'Fees', 'app.fee_payer_name', width=20,
           help='If it is not the person the learner lives with.'),
    Column('fee_payer_can_afford', 'Fee payer can afford the fees', 'Fees',
           'app.fee_payer_can_afford', 'bool', choices=YES_NO, width=10),
    Column('siblings_at_ucs', 'Siblings at UCS', 'Fees', 'app.siblings_at_ucs', 'int', width=10,
           help='Number of brothers / sisters already at UCS (5% sibling discount).'),
    Column('sibling_names', "Siblings' names and grades", 'Fees', 'app.sibling_names', width=24),
    Column('smsweb_number', 'SMSWEB contact number', 'Fees', 'app.smsweb_number', 'phone', width=16,
           help='The number the school sends SMS notifications to.'),
    # -- declarations ---------------------------------------------------------
    Column('declarations', 'Declarations signed on paper', 'Declarations', 'declarations', 'bool',
           choices=YES_NO, width=12,
           help='Yes = the parent signed the terms, indemnity, codes of conduct, prospectus, fees, '
                'POPIA and acknowledgement of documents on the paper form.'),
    Column('signed_by', 'Signed by (parent name)', 'Declarations', 'app.signed_by', width=22),
    Column('extramural_participation', 'Extra-mural programme', 'Declarations',
           'app.extramural_participation', 'bool', choices=YES_NO, default=True, width=10,
           help='Thursday 13h30 – 15h00 programme. Blank = Yes.'),
    Column('media_consent', 'Media consent', 'Declarations', 'app.media_consent', 'bool',
           choices=YES_NO, default=True, width=10,
           help='Photos / video of the learner in school publications. Blank = Yes.'),
    # -- office ---------------------------------------------------------------
    Column('office_account_number', 'ACC No', 'Office use', 'app.office_account_number', width=12,
           help='The learner\'s account number on Pastel.'),
    Column('office_pastel_account', 'Account created on Pastel', 'Office use',
           'app.office_pastel_account', 'bool', choices=YES_NO, width=10),
    Column('office_smsweb', 'Added to SMSWEB', 'Office use', 'app.office_smsweb', 'bool',
           choices=YES_NO, width=10),
    Column('office_learner_profile', 'Learner profile on SA-SAMS', 'Office use',
           'app.office_learner_profile', 'bool', choices=YES_NO, width=10),
    Column('office_transfer_received', 'Transfer card received', 'Office use',
           'app.office_transfer_received', 'bool', choices=YES_NO, width=10),
    Column('office_notes', 'Office notes', 'Office use', 'app.office_notes', width=30),
]

COLUMN_BY_KEY = {c.key: c for c in COLUMNS}
SUBJECT_COLUMNS = [c for c in COLUMNS if c.kind == 'subject']
SECTIONS = list(dict.fromkeys(c.section for c in COLUMNS))

#: Three realistic example rows (skipped on import).
EXAMPLES = [
    {
        'first_name': 'EXAMPLE Thandi', 'last_name': 'Mokoena', 'gender': 'Female',
        'date_of_birth': date(2019, 4, 12), 'email': '', 'address': '12 Raleigh Street',
        'suburb': 'Yeoville', 'city': 'Johannesburg', 'province': 'Gauteng',
        'year': school.YEAR, 'grade': 1, 'class_code': '', 'is_new_learner': 'Yes',
        'status': 'Admitted', 'paid_up_to': f'{school.YEAR}-03',
        'id_document_type': 'Birth certificate', 'id_number': '1904120800085',
        'is_south_african': 'Yes', 'home_language': 'isiZulu', 'religion': 'Christian',
        'writing_hand': 'Right', 'has_father': 'Yes', 'has_mother': 'Yes',
        'lives_with': 'Both parents',
        'g1_role': 'Mother', 'g1_title': 'Mrs', 'g1_name': 'Nomsa Mokoena',
        'g1_id': '8506150800087', 'g1_cell': '082 555 0101', 'g1_email': 'nomsa.mokoena@example.com',
        'g1_address': '12 Raleigh Street, Yeoville, Johannesburg', 'g1_occupation': 'Nurse',
        'g1_employer': 'Charlotte Maxeke Hospital', 'g1_employer_phone': '011 488 4911',
        'g2_role': 'Father', 'g2_title': 'Mr', 'g2_name': 'Sipho Mokoena', 'g2_id': '8309095800083',
        'g2_cell': '083 555 0102', 'g2_email': 'sipho.mokoena@example.com',
        'g2_occupation': 'Electrician', 'g2_employer': 'City Power',
        'emergency_name': 'Grace Dlamini', 'emergency_relationship': 'Aunt',
        'emergency_cell_phone': '072 555 0103',
        'has_medical_condition': 'Yes', 'medical_conditions': 'Mild asthma',
        'medication': 'Asthma pump when needed', 'medical_aid_name': 'Discovery Health',
        'medical_aid_number': '123456789', 'medical_aid_plan': 'KeyCare',
        'medical_aid_main_member': 'Nomsa Mokoena', 'doctor_contact': 'Dr Naidoo 011 555 0199',
        'fee_payer': 'Both parents', 'fee_payer_can_afford': 'Yes', 'siblings_at_ucs': 1,
        'sibling_names': 'Lerato Mokoena (Grade 5)', 'smsweb_number': '082 555 0101',
        'declarations': 'Yes', 'signed_by': 'Nomsa Mokoena', 'office_account_number': 'MOK001',
    },
    {
        'first_name': 'EXAMPLE Lerato', 'last_name': 'Mokoena', 'gender': 'Female',
        'date_of_birth': date(2015, 8, 3), 'email': 'lerato.mokoena@example.com',
        'phone': '', 'address': '12 Raleigh Street', 'suburb': 'Yeoville', 'city': 'Johannesburg',
        'province': 'Gauteng', 'year': school.YEAR, 'grade': 5, 'is_new_learner': 'No',
        'status': 'Admitted', 'id_document_type': 'Birth certificate', 'id_number': '1508030800082',
        'is_south_african': 'Yes', 'home_language': 'isiZulu', 'writing_hand': 'Left',
        'previous_school': '', 'has_father': 'Yes', 'has_mother': 'Yes',
        'lives_with': 'Both parents',
        'g1_role': 'Mother', 'g1_title': 'Mrs', 'g1_name': 'Nomsa Mokoena',
        'g1_cell': '082 555 0101', 'g1_email': 'nomsa.mokoena@example.com',
        'fee_payer': 'Both parents', 'siblings_at_ucs': 1,
        'sibling_names': 'Thandi Mokoena (Grade 1)', 'declarations': 'Yes',
        'signed_by': 'Nomsa Mokoena', 'office_account_number': 'MOK002',
        'office_pastel_account': 'Yes',
    },
    {
        'first_name': 'EXAMPLE Kwame', 'last_name': 'Mensah', 'gender': 'Male',
        'date_of_birth': date(2010, 1, 27), 'email': 'kwame.mensah@example.com',
        'phone': '071 555 0201', 'address': '44 Frances Street', 'suburb': 'Yeoville',
        'city': 'Johannesburg', 'province': 'Gauteng', 'year': school.YEAR, 'grade': 10,
        'class_code': '10A', 'is_new_learner': 'Yes', 'status': 'Admitted',
        'fal': 'Afrikaans First Additional Language', 'maths': 'Mathematics',
        'elective1': 'Physical Sciences', 'elective2': 'Life Sciences', 'elective3': 'Accounting',
        'paid_up_to': f'{school.YEAR}-06', 'highest_grade_passed': 'Grade 9',
        'year_grade_passed': school.YEAR - 1,
        'id_document_type': 'Passport', 'passport_number': 'G1234567', 'document_country': 'Ghana',
        'document_expiry': date(school.YEAR + 4, 5, 31), 'is_south_african': 'No',
        'study_permit_number': 'SP998877', 'study_permit_expiry': date(school.YEAR + 1, 12, 31),
        'home_language': 'English', 'religion': 'Christian', 'writing_hand': 'Right',
        'previous_school': 'Accra Academy', 'previous_school_address': 'Bubuashie, Accra, Ghana',
        'has_father': 'Yes', 'has_mother': 'No', 'lives_with': 'Uncle',
        'g1_role': 'Legal guardian', 'g1_title': 'Dr', 'g1_name': 'Kofi Mensah',
        'g1_id': 'G7654321', 'g1_cell': '084 555 0202', 'g1_work': '011 555 0203',
        'g1_email': 'kofi.mensah@example.com', 'g1_address': '44 Frances Street, Yeoville',
        'g1_occupation': 'Lecturer', 'g1_employer': 'University of the Witwatersrand',
        'emergency_name': 'Ama Owusu', 'emergency_relationship': 'Family friend',
        'emergency_cell_phone': '076 555 0204', 'has_medical_condition': 'No',
        'fee_payer': 'Guardian', 'fee_payer_can_afford': 'Yes', 'siblings_at_ucs': 0,
        'smsweb_number': '084 555 0202', 'declarations': 'Yes', 'signed_by': 'Kofi Mensah',
        'media_consent': 'No', 'office_account_number': 'MEN001',
    },
]


# ---------------------------------------------------------------------------
# Parsing cells
# ---------------------------------------------------------------------------
class CellError(ValueError):
    pass


def _norm(text):
    return re.sub(r'\s+', ' ', str(text)).strip().casefold()


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    return re.sub(r'[ \t]+', ' ', str(value)).strip()


def _choice(column, value):
    raw = _text(value)
    key = _norm(raw)
    for stored, label in column.choices:
        if key in (_norm(stored), _norm(label)) or (label and key == _norm(label.split(' (')[0])):
            return stored
    if column.kind == 'bool' and key in ALIASES:
        return ALIASES[key]
    if column.key == 'gender' and key in ('m', 'f', 'boy', 'girl'):
        return 'male' if key in ('m', 'boy') else 'female'
    if column.key == 'status':
        for stored, label in Application.STATUS_CHOICES:
            if key == _norm(label):
                return stored
    if column.key == 'id_document_type':
        for stored, label in Application.ID_DOCUMENT_CHOICES:
            if key == _norm(label):
                return stored
    if column.key == 'race' and key in ('prefer not to say', 'n/a', '-'):
        return ''
    if not column.strict:
        return raw
    raise CellError(f'"{raw}" is not one of: {", ".join(column.labels())}.')


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value)
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%d/%m/%Y', '%d-%m-%Y', '%d %B %Y', '%d %b %Y'):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise CellError(f'"{raw}" is not a date — use YYYY-MM-DD.')


def _month(value):
    """The last day of the month a ``YYYY-MM`` (or a date) names."""
    if isinstance(value, (date, datetime)):
        return school.month_end(value.date() if isinstance(value, datetime) else value)
    raw = _text(value)
    for fmt in ('%Y-%m', '%Y/%m', '%m/%Y', '%B %Y', '%b %Y', '%Y-%m-%d', '%d/%m/%Y'):
        try:
            return school.month_end(datetime.strptime(raw, fmt).date())
        except ValueError:
            continue
    raise CellError(f'"{raw}" is not a month — use YYYY-MM, e.g. {school.YEAR}-03.')


def _int(value):
    if isinstance(value, bool):
        raise CellError('expected a whole number.')
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    raw = _text(value)
    if re.fullmatch(r'\d+', raw):
        return int(raw)
    raise CellError(f'"{raw}" is not a whole number.')


def _phone(value):
    if isinstance(value, (int, float)):
        digits = str(int(value))
        # Excel drops the leading 0 of a local number typed as a number.
        return '0' + digits if len(digits) == 9 else digits
    return _text(value)


EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')


def _email(value):
    raw = _text(value).lower()
    if not EMAIL_RE.match(raw):
        raise CellError(f'"{raw}" is not a valid e-mail address.')
    return raw


def _grade(value):
    raw = _text(value)
    match = re.fullmatch(r'(?:grade\s*|gr\s*)?0*(\d{1,2})', raw, re.I)
    if match and 1 <= int(match.group(1)) <= 12:
        return int(match.group(1))
    raise CellError(f'"{raw}" is not a grade — use a number from 1 to 12.')


def parse_cell(column, value):
    """The Python value of one cell (``None`` when blank). Raises :class:`CellError`."""
    if _is_blank(value):
        return None
    kind = column.kind
    if kind in ('choice', 'bool'):
        return _choice(column, value)
    if kind == 'date':
        return _date(value)
    if kind == 'month':
        return _month(value)
    if kind == 'int':
        return _int(value)
    if kind == 'phone':
        return _phone(value)
    if kind == 'email':
        return _email(value)
    if kind == 'grade':
        return _grade(value)
    return _text(value)        # text, subject (resolved against the grade later)


def format_cell(column, value):
    """The cell value written for ``value`` (the inverse of :func:`parse_cell`)."""
    if value is None or value == '':
        return None
    if column.kind == 'bool':
        return 'Yes' if value else 'No'
    if column.kind == 'choice':
        for stored, label in column.choices:
            if stored == value:
                return label or None
        return str(value)
    if column.kind == 'month':
        return f'{value:%Y-%m}'
    return value


# ---------------------------------------------------------------------------
# The workbook (template and export share it)
# ---------------------------------------------------------------------------
def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style='thin', color='C9D3E0')
    return {
        'section': (Font(bold=True, color='FFFFFF', size=11),
                    PatternFill('solid', fgColor='003366'), Alignment(horizontal='left',
                                                                      vertical='center')),
        'header': (Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor=NAVY),
                   Alignment(wrap_text=True, vertical='center', horizontal='center')),
        'required': (Font(bold=True, color='1F2937'), PatternFill('solid', fgColor=GOLD),
                     Alignment(wrap_text=True, vertical='center', horizontal='center')),
        'example': (Font(italic=True, color='6B7280'), PatternFill('solid', fgColor='F3F4F6'), None),
        'border': Border(left=thin, right=thin, top=thin, bottom=thin),
    }


def _apply(cell, style):
    font, fill, alignment = style
    cell.font = font
    cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment


HEADER_ROW = 2      # row 1 carries the section bands
#: Columns formatted as text so Excel keeps leading zeros and long numbers.
TEXT_KEYS = {'id_number', 'g1_id', 'g2_id', 'passport_number', 'class_code', 'medical_aid_number',
             'medical_aid_main_member_id', 'office_account_number', 'study_permit_number',
             'permit_number'}
FIRST_DATA_ROW = 3


def build_workbook(rows=(), *, examples=True, title='UCS learner import'):
    """A workbook in the import format. ``rows`` are ``{column key: cell value}``
    dicts (already formatted, see :func:`format_cell`)."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    styles = _styles()
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_LEARNERS
    wb.properties.title = title
    wb.properties.creator = school.SCHOOL['name']

    # Section bands (row 1) and headers (row 2).
    start = 1
    for index, column in enumerate(COLUMNS, 1):
        last = index == len(COLUMNS) or COLUMNS[index].section != column.section
        if last:
            if index > start:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=index)
            cell = ws.cell(row=1, column=start, value=column.section.upper())
            _apply(cell, styles['section'])
            start = index + 1
        cell = ws.cell(row=HEADER_ROW, column=index, value=column.header)
        _apply(cell, styles['required' if column.required else 'header'])
        cell.border = styles['border']
        note = column.help or ''
        if column.required:
            note = ('REQUIRED. ' + note).strip()
        if column.choices and column.kind != 'grade':
            note += ('\n' if note else '') + ('Suggested: ' if not column.strict else 'One of: ') + \
                ', '.join(column.labels())
        if note:
            cell.comment = Comment(note, 'UCS', width=320, height=140)
        letter = get_column_letter(index)
        ws.column_dimensions[letter].width = max(column.width, 9)
        if column.kind in ('phone', 'month') or column.key in TEXT_KEYS:
            for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + TEMPLATE_ROWS):
                ws.cell(row=r, column=index).number_format = '@'
        elif column.kind == 'date':
            for r in range(FIRST_DATA_ROW, FIRST_DATA_ROW + TEMPLATE_ROWS):
                ws.cell(row=r, column=index).number_format = 'yyyy-mm-dd'
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[HEADER_ROW].height = 48
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=3)     # names stay in view

    # Hidden "Lists" sheet feeding the drop-downs.
    lists = wb.create_sheet(SHEET_LISTS)
    list_col = 0
    for index, column in enumerate(COLUMNS, 1):
        labels = column.labels()
        if not labels:
            continue
        list_col += 1
        lists.cell(row=1, column=list_col, value=column.header)
        for r, label in enumerate(labels, 2):
            lists.cell(row=r, column=list_col, value=label)
        col_letter = get_column_letter(list_col)
        dv = DataValidation(
            type='list', formula1=f"={SHEET_LISTS}!${col_letter}$2:${col_letter}${len(labels) + 1}",
            allow_blank=True, showErrorMessage=column.strict,
            errorTitle=column.header, error='Choose a value from the list.',
            promptTitle=column.header[:32], prompt=(column.help or 'Choose from the list.')[:250],
            showInputMessage=bool(column.help))
        letter = get_column_letter(index)
        dv.add(f'{letter}{FIRST_DATA_ROW}:{letter}{FIRST_DATA_ROW + TEMPLATE_ROWS - 1}')
        ws.add_data_validation(dv)
    lists.sheet_state = 'hidden'

    # Data.
    r = FIRST_DATA_ROW
    if examples:
        for example in EXAMPLES:
            for index, column in enumerate(COLUMNS, 1):
                cell = ws.cell(row=r, column=index, value=example.get(column.key) or None)
                _apply(cell, styles['example'])
            r += 1
    for row in rows:
        for index, column in enumerate(COLUMNS, 1):
            value = row.get(column.key)
            if value is not None and value != '':
                ws.cell(row=r, column=index, value=value)
        r += 1

    _instructions_sheet(wb, styles)
    _subjects_sheet(wb, styles)
    wb.move_sheet(SHEET_LISTS, offset=len(wb.sheetnames))
    return wb


def _instructions_sheet(wb, styles):
    from openpyxl.styles import Alignment, Font
    ws = wb.create_sheet(SHEET_INSTRUCTIONS)
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 46
    ws.column_dimensions['E'].width = 70
    ws['A1'] = f'{school.SCHOOL["name"]} — learner import'
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    intro = [
        'One row per learner per school year, on the "Learners" sheet. Keep the header rows as '
        'they are; columns may be left blank unless marked required (gold headers).',
        'The three grey EXAMPLE rows show how to fill the sheet in. Rows whose first name starts '
        'with "EXAMPLE" are never imported — overwrite or delete them.',
        'Re-importing is safe: a learner is matched on the login e-mail, else the ID / birth '
        'certificate number, else first name + surname + date of birth, and updated. Blank cells '
        'keep what is already recorded.',
        f'No learner e-mail? A login like firstname.surname.1@{GENERATED_EMAIL_DOMAIN} is '
        'generated (the import report lists them). Keep it in the sheet when you re-import.',
        'Subjects stay LOCKED until fees are paid, unless "Fees paid up to (month)" is filled in.',
        'Grade 10 – 12: fill in FAL, Maths and three electives (see the "Subjects" sheet), or '
        'leave all five blank for the default subjects.',
        'Upload the file on Admissions → Import learners. You first see a dry-run preview of '
        'every row with its errors and warnings; nothing is saved until you confirm.',
        'An export (Admissions → Export learners) uses exactly this format, so it can be edited '
        'and imported again.',
    ]
    for i, line in enumerate(intro, 3):
        ws.cell(row=i, column=1, value=f'• {line}')
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[i].height = 30
    r = len(intro) + 4
    for c, title in enumerate(('Column', 'Section', 'Required', 'Accepted values', 'Notes'), 1):
        _apply(ws.cell(row=r, column=c, value=title), styles['header'])
    for column in COLUMNS:
        r += 1
        if column.kind == 'date':
            accepted = 'Date, YYYY-MM-DD'
        elif column.kind == 'month':
            accepted = 'Month, YYYY-MM'
        elif column.kind == 'int':
            accepted = 'Whole number'
        elif column.kind == 'email':
            accepted = 'E-mail address'
        elif column.kind == 'phone':
            accepted = 'Phone number (as text)'
        elif column.kind == 'grade':
            accepted = '1 – 12'
        elif column.choices:
            accepted = (('Suggested: ' if not column.strict else '') + ', '.join(column.labels()))
        else:
            accepted = 'Text'
        notes = column.help
        if column.default is not None and 'Blank =' not in notes:
            notes = (notes + ' ' if notes else '') + \
                f'Blank = {format_cell(column, column.default)} (new learners).'
        values = (column.header, column.section, 'Yes' if column.required else '', accepted, notes)
        for c, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        if column.required:
            ws.cell(row=r, column=1).font = Font(bold=True)


def _subjects_sheet(wb, styles):
    from openpyxl.styles import Alignment
    ws = wb.create_sheet(SHEET_SUBJECTS)
    names = {s['code']: s['name'] for s in school.SUBJECTS}
    group_labels = {key: label for key, (label, _pick) in school.SUBJECT_GROUPS.items()}
    for c, (title, width) in enumerate((('Grade', 8), ('Subject', 38), ('Code', 14),
                                        ('Type', 52)), 1):
        _apply(ws.cell(row=1, column=c, value=title), styles['header'])
        ws.column_dimensions['ABCD'[c - 1]].width = width
    r = 1
    for grade in school.GRADES:
        for code, group, _note in school.subjects_for(grade):
            r += 1
            for c, value in enumerate((grade, names[code], code,
                                       group_labels.get(group, 'Compulsory')), 1):
                ws.cell(row=r, column=c, value=value).alignment = Alignment(vertical='top')
    ws.freeze_panes = 'A2'


def write_template(path):
    """Write the import template (with its EXAMPLE rows) to ``path``."""
    wb = build_workbook(examples=True)
    wb.save(path)
    return path


def template_bytes():
    buffer = io.BytesIO()
    build_workbook(examples=True).save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Reading a workbook
# ---------------------------------------------------------------------------
class WorkbookError(ValueError):
    """The file cannot be read as a learner sheet at all."""


def _header_key(text):
    return _norm(str(text).rstrip('*').strip())


HEADER_LOOKUP = {_header_key(c.header): c for c in COLUMNS}
HEADER_LOOKUP.update({_header_key(c.key): c for c in COLUMNS})


def read_rows(file):
    """``(columns_found, [(row number, {key: value})], warnings)`` from a workbook.

    ``file`` is a path, bytes or a file object."""
    from openpyxl import load_workbook
    if isinstance(file, (bytes, bytearray)):
        file = io.BytesIO(file)
    try:
        wb = load_workbook(file, data_only=True, read_only=True)
    except Exception as exc:
        raise WorkbookError(f'This file could not be opened as an Excel workbook ({exc}).')
    ws = wb[SHEET_LEARNERS] if SHEET_LEARNERS in wb.sheetnames else wb.worksheets[0]
    first = _header_key(COLUMN_BY_KEY['first_name'].header)
    header_row, mapping, warnings = None, {}, []
    rows = ws.iter_rows(values_only=True)
    number = 0
    for values in rows:
        number += 1
        keys = [_header_key(v) if v is not None else '' for v in values]
        if first in keys or 'first_name' in keys:
            header_row = number
            for index, key in enumerate(keys):
                column = HEADER_LOOKUP.get(key)
                if column is not None:
                    mapping[index] = column
                elif key:
                    warnings.append(f'Column "{values[index]}" is not part of the format and was '
                                    'ignored.')
            break
        if number >= 10:
            break
    if header_row is None:
        wb.close()
        raise WorkbookError('No header row found: the sheet must have a "First name(s)" column. '
                            'Start from the UCS learner import template.')
    found = {c.key for c in mapping.values()}
    missing = [c.header for c in COLUMNS if c.required and c.key not in found]
    if missing:
        wb.close()
        raise WorkbookError('Required column(s) missing: ' + ', '.join(missing) + '.')
    data = []
    for values in rows:
        number += 1
        row = {column.key: values[index] for index, column in mapping.items()
               if index < len(values)}
        if all(_is_blank(v) for v in row.values()):
            continue
        data.append((number, row))
    wb.close()
    return found, data, warnings


# ---------------------------------------------------------------------------
# The import report
# ---------------------------------------------------------------------------
@dataclass
class RowResult:
    row: int
    name: str = ''
    email: str = ''
    grade: object = None
    action: str = ''            # created / updated / skipped / error
    errors: list = field(default_factory=list)      # [(column header, message)]
    warnings: list = field(default_factory=list)    # [(column header, message)]
    notes: list = field(default_factory=list)       # what was done

    def error(self, column, message):
        self.errors.append((column, message))

    def warn(self, column, message):
        self.warnings.append((column, message))

    def as_dict(self):
        return {'row': self.row, 'name': self.name, 'email': self.email, 'grade': self.grade,
                'action': self.action, 'errors': self.errors, 'warnings': self.warnings,
                'notes': self.notes}


@dataclass
class ImportReport:
    dry_run: bool = True
    rows: list = field(default_factory=list)
    warnings: list = field(default_factory=list)    # sheet-level
    fatal: str = ''

    def count(self, action):
        return sum(1 for r in self.rows if r.action == action)

    @property
    def counts(self):
        return {a: self.count(a) for a in ('created', 'updated', 'skipped', 'error')}

    @property
    def ok(self):
        return not self.fatal and not self.count('error')

    def summary(self):
        if self.fatal:
            return self.fatal
        c = self.counts
        verb = 'would be' if self.dry_run else ''
        return (f"{c['created']} learner(s) {verb} created, {c['updated']} {verb} updated, "
                f"{c['skipped']} skipped, {c['error']} with errors.").replace('  ', ' ')

    def as_dict(self):
        return {'dry_run': self.dry_run, 'fatal': self.fatal, 'warnings': self.warnings,
                'counts': self.counts, 'rows': [r.as_dict() for r in self.rows]}


# ---------------------------------------------------------------------------
# Importing
# ---------------------------------------------------------------------------
def _slug(text):
    text = unicodedata.normalize('NFKD', text or '').encode('ascii', 'ignore').decode()
    return re.sub(r'[^a-z0-9]+', '', text.lower().split(' ')[0]) or 'learner'


def generate_email(first_name, last_name, taken=()):
    """``firstname.surname.N@learners.ucs.local``, unique among users."""
    User = get_user_model()
    base = f'{_slug(first_name)}.{_slug(last_name)}'
    n = 1
    while True:
        email = f'{base}.{n}@{GENERATED_EMAIL_DOMAIN}'
        if email not in taken and not User.objects.filter(
                Q(username__iexact=email) | Q(email__iexact=email)).exists():
            return email
        n += 1


def is_generated_email(email):
    return bool(email) and email.lower().endswith('@' + GENERATED_EMAIL_DOMAIN)


def _ucs_programmes():
    from apps.learning.models import Programme
    return {p.grade: p for p in Programme.objects.filter(
        institution__code=school.SCHOOL['code'], grade__isnull=False, is_active=True)}


def _resolve_subject(programme, column, raw):
    """The offering ``raw`` (a subject name or code) names in ``programme``'s ``column.group``."""
    key = _norm(raw)
    for offering in programme.modules.filter(is_active=True).select_related('module'):
        if offering.subject_group != column.group:
            continue
        if key in (_norm(offering.code), _norm(offering.display_name), _norm(offering.module.name)):
            return offering
    allowed = ', '.join(o.display_name for o in programme.modules.filter(
        is_active=True, subject_group=column.group))
    raise CellError(f'"{raw}" is not offered in {programme.display_name} for this choice '
                    f'(choose from: {allowed}).')


class _Importer:
    def __init__(self, *, dry_run, create_parent_accounts, send_invites, user):
        self.dry_run = dry_run
        self.create_parent_accounts = create_parent_accounts
        self.send_invites = send_invites
        self.user = user if getattr(user, 'is_authenticated', False) else None
        self.programmes = _ucs_programmes()
        self.seen = {}           # learner pk → first row number
        self.generated = set()

    # -- validation -----------------------------------------------------------
    def clean(self, number, raw, result):
        values = {}
        for key, value in raw.items():
            column = COLUMN_BY_KEY[key]
            try:
                values[key] = parse_cell(column, value)
            except CellError as exc:
                result.error(column.header, str(exc))
        result.name = f"{values.get('first_name') or ''} {values.get('last_name') or ''}".strip()
        result.email = values.get('email') or ''
        result.grade = values.get('grade')
        for column in COLUMNS:
            if column.required and values.get(column.key) is None and not any(
                    h == column.header for h, _ in result.errors):
                result.error(column.header, 'Required.')

        year = values.get('year')
        if year is not None and not 2000 <= year <= 2100:
            result.error(COLUMN_BY_KEY['year'].header, f'{year} is not a school year.')

        grade = values.get('grade')
        programme = self.programmes.get(grade) if grade else None
        if grade and programme is None:
            result.error(COLUMN_BY_KEY['grade'].header,
                         f'Grade {grade} is not set up on the platform (run seed_school_structure).')
        values['_programme'] = programme

        # Grade 10 – 12 subject choices.
        given = {c.key: values.get(c.key) for c in SUBJECT_COLUMNS if values.get(c.key)}
        values['_subjects'] = None
        if given and programme is not None:
            if grade < 10:
                for key in given:
                    result.warn(COLUMN_BY_KEY[key].header,
                                f'Subject choices apply to Grade 10 – 12 only; ignored for '
                                f'Grade {grade}.')
            else:
                chosen = {}
                for column in SUBJECT_COLUMNS:
                    raw_value = values.get(column.key)
                    if not raw_value:
                        result.error(column.header, 'Fill in all five Grade 10 – 12 subject '
                                                    'choices, or leave all of them blank.')
                        continue
                    try:
                        chosen[column.key] = _resolve_subject(programme, column, raw_value)
                    except CellError as exc:
                        result.error(column.header, str(exc))
                electives = [chosen[k] for k in ('elective1', 'elective2', 'elective3')
                             if k in chosen]
                if len({o.pk for o in electives}) != len(electives):
                    result.error(COLUMN_BY_KEY['elective3'].header,
                                 'The three electives must be different subjects.')
                values['_subjects'] = list(chosen.values())

        # Parents.
        roles = {}
        for n in (1, 2):
            p = f'g{n}'
            block = {c.key: values.get(c.key) for c in COLUMNS if c.obj == p}
            filled = {k: v for k, v in block.items() if v not in (None, '')}
            name_header = COLUMN_BY_KEY[f'{p}_name'].header
            if not filled:
                continue
            if not values.get(f'{p}_name'):
                result.error(name_header, f'Parent {n} details are filled in but the name is blank.')
                continue
            role = values.get(f'{p}_role')
            if not role:
                result.error(COLUMN_BY_KEY[f'{p}_role'].header,
                             f'Choose Father, Mother or Legal guardian for parent {n}.')
                continue
            if role in roles:
                result.error(COLUMN_BY_KEY[f'{p}_role'].header,
                             f'Parent 1 and parent 2 are both "{dict(ROLES)[role]}".')
            roles[role] = n
        if values.get('email') and any(values.get(f'g{n}_email') == values['email']
                                       for n in (1, 2)):
            result.warn(COLUMN_BY_KEY['email'].header,
                        'The learner e-mail is the same as a parent\'s; the learner account uses it '
                        'and no parent account is made for that parent.')
        return values

    # -- matching -------------------------------------------------------------
    def find_learner(self, values, result):
        from apps.accounts.models import Person
        User = get_user_model()
        email = values.get('email')
        if email:
            user = User.objects.filter(Q(email__iexact=email) | Q(username__iexact=email)).first()
            if user is not None:
                person = Person.objects.filter(user=user).first()
                if person is not None and person.user_type != 'student':
                    raise CellError(f'{email} belongs to an existing {person.user_type} account, '
                                    'not a learner.')
                return person
        id_number = values.get('id_number')
        if id_number:
            application = (Application.objects.filter(id_number__iexact=id_number)
                           .select_related('person').order_by('-year').first())
            if application is not None:
                return application.person
        if values.get('first_name') and values.get('last_name') and values.get('date_of_birth'):
            return (Person.objects.filter(
                user_type='student', first_name__iexact=values['first_name'],
                last_name__iexact=values['last_name'], date_of_birth=values['date_of_birth'])
                .order_by('pk').first())
        return None

    # -- applying a row -------------------------------------------------------
    def apply(self, values, result):
        from apps.accounts.models import Person, PersonContact
        from core.seed_builders import verify_email
        User = get_user_model()

        created = False
        try:
            person = self.find_learner(values, result)
        except CellError as exc:
            result.error(COLUMN_BY_KEY['email'].header, str(exc))
            return
        if person is not None and person.pk in self.seen:
            result.warn('', f'Same learner as row {self.seen[person.pk]}; this row updates it again.')

        email = values.get('email')
        if person is None:
            if not email:
                email = generate_email(values['first_name'], values['last_name'], self.generated)
                self.generated.add(email)
                result.warn(COLUMN_BY_KEY['email'].header,
                            f'No e-mail: login {email} generated (it cannot receive mail).')
            user = User.objects.create_user(username=email, email=email)
            user.set_unusable_password()
            user.first_name = values['first_name'][:150]
            user.last_name = values['last_name'][:150]
            user.save()
            person = Person.objects.filter(user=user).first() or Person.objects.create(user=user)
            if not is_generated_email(email):
                verify_email(user)
            created = True
        else:
            user = person.user
            if email and email != (user.email or '').lower():
                if User.objects.filter(Q(email__iexact=email) | Q(username__iexact=email)).exclude(
                        pk=user.pk).exists():
                    result.error(COLUMN_BY_KEY['email'].header,
                                 f'{email} is already used by another account.')
                    return
                result.warn(COLUMN_BY_KEY['email'].header,
                            f'Login e-mail changed from {user.email} to {email}.')
                user.email = user.username = email
                user.save(update_fields=['email', 'username'])
                if not is_generated_email(email):
                    verify_email(user)
        self.seen.setdefault(person.pk, result.row)
        result.email = user.email
        result.action = 'created' if created else 'updated'

        # Person + contact
        person.user_type = 'student'
        person.registered = True
        person.profile_status = True
        for column in COLUMNS:
            if column.obj == 'person' and values.get(column.key) is not None:
                setattr(person, column.attr, values[column.key])
        programme = values['_programme']
        person.enrolled_class = programme.display_name[:80]
        person.save()
        contact, _ = PersonContact.objects.get_or_create(person=person)
        for column in COLUMNS:
            if column.obj == 'contact' and values.get(column.key) is not None:
                setattr(contact, column.attr, values[column.key])
        if values.get('phone') is not None:
            contact.primary_phone = values['phone']
        contact.save()

        # Application
        year = values.get('year') or school.YEAR
        application = Application.objects.filter(person=person, year=year).first()
        new_application = application is None
        if new_application:
            application = Application(person=person, year=year)
        for column in COLUMNS:
            if column.obj != 'app' or column.attr == 'year':
                continue
            value = values.get(column.key)
            if value is None and new_application and column.default is not None:
                value = column.default
            if value is not None:
                setattr(application, column.attr, value)
        application.programme = programme
        if values.get('gender'):
            application.gender = values['gender']
        if values.get('declarations') is not None:
            for name in REQUIRED_DECLARATIONS:
                setattr(application, name, values['declarations'])
            if values['declarations'] and not application.signed_at:
                application.signed_at = timezone.now()
            if values['declarations'] and not application.signed_by:
                application.signed_by = (values.get('g1_name') or '')[:120]
        if application.status not in (Application.STATUS_DRAFT,) and not application.submitted_at:
            application.submitted_at = timezone.now()
        if application.status == Application.STATUS_ADMITTED and not application.decided_at:
            application.decided_at = timezone.now()
            application.decided_by = self.user
        if not application.office_notes and new_application:
            application.office_notes = 'Imported from the learner spreadsheet.'
        application.save()
        if not new_application:
            result.notes.append(f'Application {year} updated.')
        else:
            result.notes.append(f'Application {year} created ({application.get_status_display()}).')

        # Guardians
        guardians = []
        for n in (1, 2):
            p = f'g{n}'
            if not values.get(f'{p}_name'):
                continue
            guardian = (application.guardians.filter(role=values[f'{p}_role']).first()
                        or Guardian(application=application, role=values[f'{p}_role']))
            for column in COLUMNS:
                if column.obj == p and values.get(column.key) is not None:
                    setattr(guardian, column.attr, values[column.key])
            guardian.save()
            guardians.append(guardian)

        self.enrol(person, application, values, result)
        self.parents(person, application, guardians, result)

    def enrol(self, person, application, values, result):
        from apps.learning.models import Cohort, ModuleEnrolment, ProgrammeEnrolment
        from core.academic_spine import default_subjects, enrol_student

        programme = values['_programme']
        year = application.year
        code = values.get('class_code')
        existing = ProgrammeEnrolment.objects.filter(person=person, programme=programme).first()
        if code is None and existing is not None and existing.cohort_id:
            cohort = existing.cohort
        else:
            code = code or str(year)
            cohort = Cohort.objects.filter(programme=programme, code__iexact=code).first()
            if cohort is None:
                cohort = Cohort.objects.create(
                    programme=programme, code=code[:30],
                    name=f'{programme.display_name} · {code}'[:120])
                result.warn(COLUMN_BY_KEY['class_code'].header,
                            f'Class "{code}" did not exist in {programme.display_name}; created.')

        offerings = values.get('_subjects')
        compulsory = list(programme.modules.filter(is_active=True, subject_group=''))
        if offerings is not None:
            offerings = compulsory + offerings
        elif (programme.grade or 0) >= 10:
            held = list(programme.modules.filter(
                is_active=True, enrolments__person=person).exclude(subject_group='').distinct())
            if held:
                offerings = compulsory + held
            else:
                offerings = default_subjects(programme)
                result.warn(COLUMN_BY_KEY['fal'].header,
                            'No subject choices given: default subjects assigned ('
                            + ', '.join(o.display_name for o in offerings if o.subject_group)
                            + ').')
        else:
            offerings = default_subjects(programme)

        enrolment, modules = enrol_student(person, programme, cohort=cohort, activate=False,
                                           offerings=offerings)
        active = application.status not in (Application.STATUS_DECLINED,
                                             Application.STATUS_WITHDRAWN)
        if enrolment.cohort_id != cohort.pk or enrolment.is_active != active:
            enrolment.cohort = cohort
            enrolment.is_active = active
            enrolment.save(update_fields=['cohort', 'is_active', 'updated_at'])
        # A learner moved to another grade is no longer active in the old one.
        moved = (ProgrammeEnrolment.objects.filter(person=person, is_active=True,
                                                   programme__institution=programme.institution)
                 .exclude(programme=programme))
        for old in moved:
            old.is_active = False
            old.save(update_fields=['is_active', 'updated_at'])
            result.warn(COLUMN_BY_KEY['grade'].header,
                        f'No longer active in {old.programme.display_name}.')
        # Subject choices dropped in this grade.
        keep = {o.pk for o in offerings}
        for dropped in ModuleEnrolment.objects.filter(
                person=person, programme_module__programme=programme).exclude(
                programme_module_id__in=keep).exclude(programme_module__subject_group=''):
            dropped.delete()

        paid = values.get('paid_up_to')
        if paid is not None:
            for module in modules:
                module.status = ModuleEnrolment.STATUS_ACTIVE
                module.paid_until = paid
                module.started_at = module.started_at or timezone.now()
                module.save(update_fields=['status', 'paid_until', 'started_at', 'updated_at'])
            result.notes.append(f'{len(modules)} subject(s) unlocked to {paid:%d %B %Y}.')
        else:
            unlocked = sum(1 for m in modules if m.status == ModuleEnrolment.STATUS_ACTIVE)
            result.notes.append(f'{len(modules)} subject(s) in {programme.display_name} · '
                                f'{cohort.code}' + (f' ({unlocked} unlocked)' if unlocked
                                                    else ' (locked until fees are paid)') + '.')

    def parents(self, person, application, guardians, result):
        from apps.accounts.models import ParentLink, Person
        from core.seed_builders import verify_email
        User = get_user_model()
        student = person.user
        with_email = [g for g in guardians if g.email and
                      g.email.lower() != (student.email or '').lower()]
        if not with_email:
            return
        if self.create_parent_accounts:
            for guardian in with_email:
                email = guardian.email.lower()
                parent = User.objects.filter(Q(email__iexact=email) | Q(username__iexact=email)).first()
                if parent is not None:
                    profile = Person.objects.filter(user=parent).first()
                    if profile is not None and profile.user_type == 'student':
                        result.warn('Parent e-mail', f'{email} is a learner account; not linked '
                                                      'as a parent.')
                        continue
                    made = False
                else:
                    parent = User.objects.create_user(username=email, email=email)
                    parent.set_unusable_password()
                    first, _, last = guardian.full_name.rpartition(' ')
                    parent.first_name, parent.last_name = (first or last)[:150], (
                        last if first else '')[:150]
                    parent.save()
                    profile = Person.objects.filter(user=parent).first() or Person.objects.create(
                        user=parent)
                    profile.user_type = 'parent'
                    profile.first_name, profile.last_name = parent.first_name[:50], \
                        parent.last_name[:50]
                    title = (guardian.title or '').lower()
                    if title in {'mr', 'mrs', 'ms', 'miss', 'dr', 'prof'}:
                        profile.title = title
                    profile.phone = guardian.cell_phone or guardian.home_phone or ''
                    profile.registered = profile.profile_status = True
                    profile.child_name = f'{person.first_name} {person.last_name}'[:160]
                    profile.save()
                    verify_email(parent)
                    made = True
                if ParentLink.objects.filter(parent=parent, student=student).exists():
                    continue
                if not ParentLink.can_add_parent(student):
                    result.warn('Parent e-mail', f'{email} not linked: a learner may have at most '
                                                 f'{ParentLink.MAX_PER_STUDENT} parent accounts.')
                    continue
                ParentLink.objects.create(parent=parent, student=student,
                                          relationship=guardian.get_role_display())
                result.notes.append(('Parent account created and linked: ' if made
                                     else 'Linked to existing parent account: ') + email)
        elif self.send_invites:
            pending = [g.email for g in with_email if not g.invite_sent_at]
            if not pending:
                return
            if self.dry_run:
                result.notes.append('Parent invitation would be sent to ' + ', '.join(pending) + '.')
                return
            from .services import invite_guardians

            def send(application_pk=application.pk, invited_by=self.user):
                try:
                    invite_guardians(Application.objects.get(pk=application_pk),
                                     invited_by=invited_by)
                except Exception:  # pragma: no cover - e-mail is best-effort
                    logger.exception('bulk import: parent invitations failed')
            transaction.on_commit(send)
            result.notes.append('Parent invitation sent to ' + ', '.join(pending) + '.')

    # -- one row --------------------------------------------------------------
    def run_row(self, number, raw):
        result = RowResult(row=number)
        first = _text(raw.get('first_name') or '')
        if first.upper().startswith(EXAMPLE_PREFIX):
            result.name = f"{first} {_text(raw.get('last_name') or '')}".strip()
            result.action = 'skipped'
            result.notes.append('Example row — not imported.')
            return result
        values = self.clean(number, raw, result)
        if result.errors:
            result.action = 'error'
            return result
        try:
            with transaction.atomic():
                self.apply(values, result)
                if result.errors:
                    raise _RowRollback()
        except _RowRollback:
            result.action = 'error'
        except Exception as exc:     # one bad row never stops the rest
            logger.exception('bulk import: row %s failed', number)
            result.action = 'error'
            result.error('', f'Could not be saved: {exc}')
        return result


class _RowRollback(Exception):
    pass


def import_workbook(file, dry_run=True, create_parent_accounts=False, send_invites=False,
                    user=None):
    """Import learners from an xlsx (path, bytes or file object).

    Returns an :class:`ImportReport`. With ``dry_run`` (the default) every row
    is processed inside a transaction that is then rolled back — the report is
    what a real import would do, and nothing is saved or sent."""
    report = ImportReport(dry_run=dry_run)
    try:
        _found, rows, warnings = read_rows(file)
    except WorkbookError as exc:
        report.fatal = str(exc)
        return report
    report.warnings = warnings
    if not rows:
        report.fatal = 'The sheet has no learner rows.'
        return report
    importer = _Importer(dry_run=dry_run, create_parent_accounts=create_parent_accounts,
                         send_invites=send_invites, user=user)
    if not importer.programmes:
        report.fatal = 'The grades are not set up yet — run seed_school_structure first.'
        return report

    if dry_run:
        with transaction.atomic():
            for number, raw in rows:
                report.rows.append(importer.run_row(number, raw))
            transaction.set_rollback(True)
    else:
        for number, raw in rows:
            report.rows.append(importer.run_row(number, raw))
    return report


# ---------------------------------------------------------------------------
# Exporting
# ---------------------------------------------------------------------------
def _row_for(application=None, person=None, enrolment=None, year=None):
    """``{column key: cell value}`` for one learner (and their application)."""
    from apps.accounts.models import PersonContact
    from apps.learning.models import ModuleEnrolment

    person = person or application.person
    contact = PersonContact.objects.filter(person=person).first()
    programme = application.programme if application and application.programme_id else (
        enrolment.programme if enrolment else None)
    # Parent 1 / parent 2 in the order they were recorded (as in the sheet).
    guardians = sorted(application.guardians.all(), key=lambda g: g.pk) if application else []
    objects = {'user': person.user, 'person': person, 'contact': contact, 'app': application,
               'g1': guardians[0] if guardians else None,
               'g2': guardians[1] if len(guardians) > 1 else None}

    row = {}
    for column in COLUMNS:
        obj = objects.get(column.obj)
        value = getattr(obj, column.attr, None) if obj is not None and column.attr else None
        row[column.key] = format_cell(column, value)
    if application is None:
        row['year'] = year or school.YEAR
    if row.get('phone') is None and contact is not None and contact.primary_phone:
        row['phone'] = contact.primary_phone
    row['grade'] = programme.grade if programme else None

    if programme is not None:
        from apps.learning.models import ProgrammeEnrolment
        enrolment = enrolment or ProgrammeEnrolment.objects.filter(
            person=person, programme=programme).select_related('cohort').first()
        if enrolment is not None and enrolment.cohort_id:
            row['class_code'] = enrolment.cohort.code
        modules = list(ModuleEnrolment.objects.filter(
            person=person, programme_module__programme=programme)
            .select_related('programme_module__module').order_by('programme_module__order'))
        if (programme.grade or 0) >= 10:
            by_group = {}
            for m in modules:
                if m.programme_module.subject_group:
                    by_group.setdefault(m.programme_module.subject_group, []).append(
                        m.programme_module.display_name)
            for column in SUBJECT_COLUMNS:
                picks = by_group.get(column.group, [])
                index = int(column.key[-1]) - 1 if column.key.startswith('elective') else 0
                row[column.key] = picks[index] if index < len(picks) else None
        paid = [m.paid_until for m in modules
                if m.status == ModuleEnrolment.STATUS_ACTIVE and m.paid_until]
        row['paid_up_to'] = f'{max(paid):%Y-%m}' if paid else None
    if application is not None:
        row['declarations'] = 'Yes' if all(
            getattr(application, n) for n in REQUIRED_DECLARATIONS) else 'No'
    return row


def export_rows(queryset=None, *, year=None, grade=None):
    """The export rows: every application of ``year`` (default the school year)
    and, when no ``queryset`` is given, enrolled learners with no application
    that year (so a register seeded without applications exports too)."""
    from apps.learning.models import ProgrammeEnrolment

    if queryset is None:
        year = year or school.YEAR
        queryset = Application.objects.filter(year=year)
        if grade:
            queryset = queryset.filter(programme__grade=grade)
        extra = (ProgrammeEnrolment.objects
                 .filter(is_active=True, person__user_type='student',
                         programme__institution__code=school.SCHOOL['code'])
                 .exclude(person__applications__year=year)
                 .select_related('person__user', 'programme', 'cohort'))
        if grade:
            extra = extra.filter(programme__grade=grade)
    else:
        extra = ProgrammeEnrolment.objects.none()
    queryset = (queryset.select_related('person__user', 'programme')
                .prefetch_related('guardians')
                .order_by('programme__grade', 'person__last_name', 'person__first_name', 'pk'))
    rows = [_row_for(application=a) for a in queryset]
    seen = set()
    for enrolment in extra.order_by('programme__grade', 'person__last_name', 'person__first_name'):
        if enrolment.person_id in seen:
            continue
        seen.add(enrolment.person_id)
        rows.append(_row_for(person=enrolment.person, enrolment=enrolment, year=year))
    return rows


def export_workbook(queryset=None, *, year=None, grade=None, rows=None):
    """The learners as xlsx bytes, in exactly the import format. ``queryset``
    (of Applications) or the ``year`` / ``grade`` filters choose the learners;
    ``rows`` from :func:`export_rows` may be passed instead."""
    if rows is None:
        rows = export_rows(queryset, year=year, grade=grade)
    wb = build_workbook(rows, examples=False, title='UCS learner export')
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
