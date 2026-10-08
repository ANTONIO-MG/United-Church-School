"""SA-SAMS export: workbook layout, code mapping, marks, attendance, data
quality, staff-only pages and the learner-numbers grid."""
import datetime
import io
import os
import tempfile
import zipfile
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.admissions import promotion
from apps.admissions.models import Application, Guardian
from apps.attendance.models import AttendanceMark, DailyRegister
from apps.learning.models import Programme
from apps.reports.models import TermResult
from apps.reports.terms import term_dates
from core import academic_spine

from . import exports

User = get_user_model()
YEAR, TERM = 2026, 3


def person(email, kind='student', **fields):
    user = User.objects.create_user(username=email, email=email, password='x-Pass-1234')
    p = user.profile
    p.user_type, p.registered, p.profile_status = kind, True, True
    p.first_name = email.split('@')[0].title()
    for key, value in fields.items():
        setattr(p, key, value)
    p.save()
    return p


def sa_id(born, female=True, citizen=True):
    """A valid 13-digit SA ID number for ``born``."""
    body = f'{born:%y%m%d}{"0123" if female else "5123"}{"0" if citizen else "1"}8'
    for check in range(10):
        if exports.luhn_ok(body + str(check)):
            return body + str(check)
    raise AssertionError('no check digit')


@override_settings(SCHOOL_EMIS_NUMBER='700123456',
                   EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class SasamsExportTests(TestCase):

    @classmethod
    def setUpTestData(cls):
        academic_spine.seed(verbose=False, calendar=False)

    def setUp(self):
        self.g5 = Programme.objects.get(institution__code='UCS', grade=5)
        self.cohort = promotion.class_for(self.g5, YEAR)
        self.good_dob = datetime.date(2015, 3, 14)
        self.good = self.learner('thandi@example.com', 'Mokoena', self.good_dob,
                                 id_number=sa_id(self.good_dob), race='african',
                                 home_language='Zulu', writing_hand='left',
                                 admission_number='A0001', lurits_number='123456789')
        self.bad = self.learner('sipho@example.com', 'Dlamini', datetime.date(2015, 6, 1),
                                gender='male', id_number='1506015123089', race='',
                                home_language='')
        self.office = person('office@example.com', 'staff')

    def learner(self, email, surname, dob, gender='female', admission_number='', lurits_number='',
                **app_fields):
        learner = person(email, last_name=surname, date_of_birth=dob, gender=gender,
                         admission_number=admission_number, lurits_number=lurits_number)
        academic_spine.enrol_student(learner, self.g5, cohort=self.cohort, activate=False)
        app = Application.objects.create(person=learner, year=YEAR, programme=self.g5,
                                         status=Application.STATUS_ADMITTED, gender=gender,
                                         id_document_type=Application.DOC_BIRTH_CERT, **app_fields)
        Guardian.objects.create(application=app, role='mother', full_name=f'Grace {surname}',
                                cell_phone='0820000000', id_number='8001010000000')
        return learner

    # -- codes & helpers -------------------------------------------------------
    def test_code_mapping(self):
        self.assertEqual(exports.gender_code('female'), 'F')
        self.assertEqual(exports.gender_code('male'), 'M')
        self.assertEqual(exports.gender_code('other'), '')
        self.assertEqual(exports.population_group('coloured'), ('Coloured', 2))
        self.assertEqual(exports.language_name('zulu'), 'isiZulu')
        self.assertEqual(exports.language_name('Northern Sotho'), 'Sepedi')
        self.assertEqual(exports.language_name('Shona'), 'Other')
        self.assertEqual(exports.check_sa_id(sa_id(self.good_dob), self.good_dob, 'female'), [])

    def test_bad_id_number_flagged(self):
        problems = exports.check_sa_id('1506015123089', datetime.date(2015, 6, 1), 'male')
        self.assertTrue(any('Luhn' in p for p in problems))
        problems = exports.check_sa_id(sa_id(self.good_dob), datetime.date(2015, 3, 15), 'female')
        self.assertTrue(any('does not match' in p for p in problems))
        self.assertTrue(exports.check_sa_id('12345', None))

        issues = exports.data_quality(exports.collect(YEAR, 5))
        bad = [i for i in issues if i['person_id'] == self.bad.pk]
        fields = {i['Field'] for i in bad}
        self.assertIn('Learner ID Number', fields)
        self.assertIn('Admission Number', fields)
        self.assertIn('LURITS Number', fields)
        self.assertIn('Learner Home Language', fields)
        self.assertIn('Learner Ethnic Group', fields)
        good_errors = [i for i in issues if i['person_id'] == self.good.pk and i['Severity'] == 'Error']
        self.assertEqual(good_errors, [])

    def test_duplicate_admission_number_flagged(self):
        self.bad.admission_number = 'a0001'
        self.bad.save()
        issues = exports.data_quality(exports.collect(YEAR, 5))
        self.assertTrue(any('used by 2 learners' in i['Problem'] for i in issues))

    # -- workbook ---------------------------------------------------------------
    def _workbook(self, **kw):
        return load_workbook(io.BytesIO(exports.workbook_bytes(YEAR, TERM, 5, **kw)))

    def test_workbook_sheets_columns_and_codes(self):
        wb = self._workbook()
        for name in ('Read me', 'Learners', 'Parents', 'Marks Gr05 T3', 'Marks T3 (list)',
                     'Attendance T3', 'Data quality'):
            self.assertIn(name, wb.sheetnames)
        ws = wb['Learners']
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers, [h for h, _ in exports.LEARNER_COLUMNS])
        rows = {r[headers.index('Learner Surname')]: dict(zip(headers, r))
                for r in ws.iter_rows(min_row=2, values_only=True)}
        thandi = rows['Mokoena']
        self.assertEqual(thandi['EMIS Number'], '700123456')
        self.assertEqual(thandi['Admission Number'], 'A0001')
        self.assertEqual(thandi['LURITS Number'], '123456789')
        self.assertEqual(thandi['Learner Gender'], 'F')
        self.assertEqual(thandi['Learner Birthdate'], '2015-03-14')
        self.assertEqual(thandi['Learner Ethnic Group'], 'African')
        self.assertEqual(thandi['Population Group Code'], 1)
        self.assertEqual(thandi['Learner Home Language'], 'isiZulu')
        self.assertEqual(thandi['Learner Dexterity'], 'Left')
        self.assertEqual(thandi['Learner Grade'], 5)
        self.assertEqual(thandi['Learner Register Class'], '5A')
        self.assertEqual(thandi['Citizenship'], 'South African')
        self.assertEqual(rows['Dlamini']['Learner Gender'], 'M')
        self.assertIn(rows['Dlamini']['Admission Number'], (None, ''))

        parents = wb['Parents']
        p_headers = [c.value for c in parents[1]]
        self.assertEqual(p_headers, [h for h, _ in exports.GUARDIAN_COLUMNS])
        mother = [dict(zip(p_headers, r)) for r in parents.iter_rows(min_row=2, values_only=True)
                  if r[p_headers.index('Learner Admission Number')] == 'A0001'][0]
        self.assertEqual(mother['Relationship'], 'Mother')
        self.assertEqual(mother['Parent Surname'], 'Mokoena')
        self.assertEqual(mother['Parent Full Names'], 'Grace')

        readme = [r for r in wb['Read me'].iter_rows(values_only=True)]
        self.assertTrue(any(r[0] == 'EMIS number' and r[1] == '700123456' for r in readme))
        dq = [c.value for c in wb['Data quality'][1]]
        self.assertEqual(dq, exports.DATA_QUALITY_COLUMNS)
        self.assertGreater(wb['Data quality'].max_row, 1)

    def test_only_published_marks_by_default(self):
        math = self.g5.modules.get(code='MATH')
        eng = self.g5.modules.get(code='ENG-HL')
        TermResult.objects.create(student=self.good.user, module=math, year=YEAR, term=TERM,
                                  sba_pct=Decimal('72'), term_pct=Decimal('72'), level=6,
                                  status=TermResult.STATUS_PUBLISHED)
        TermResult.objects.create(student=self.good.user, module=eng, year=YEAR, term=TERM,
                                  term_pct=Decimal('45'), level=3, status=TermResult.STATUS_DRAFT)

        def grid(wb):
            ws = wb['Marks Gr05 T3']
            headers = [c.value for c in ws[1]]
            row = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == 'A0001'][0]
            return dict(zip(headers, row))

        published = grid(self._workbook())
        self.assertEqual(published['MATH %'], 72)
        self.assertEqual(published['MATH Level'], 6)
        self.assertIn(published['ENG-HL %'], (None, ''))
        listed = list(self._workbook()['Marks T3 (list)'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(listed), 1)

        with_drafts = grid(self._workbook(include_drafts=True))
        self.assertEqual(with_drafts['ENG-HL %'], 45)
        self.assertEqual(with_drafts['ENG-HL Level'], 3)

    def test_attendance_summary(self):
        start, _end = term_dates(YEAR, TERM)
        for offset, status in enumerate(['present', 'absent', 'late', 'excused', 'present']):
            register = DailyRegister.objects.create(cohort=self.cohort,
                                                    date=start + datetime.timedelta(days=offset))
            AttendanceMark.objects.create(register=register, learner=self.good, status=status)
        ws = self._workbook()['Attendance T3']
        headers = [c.value for c in ws[1]]
        row = dict(zip(headers, [r for r in ws.iter_rows(min_row=2, values_only=True)
                                 if r[headers.index('Admission Number')] == 'A0001'][0]))
        self.assertEqual(row['School Days Recorded'], 5)
        self.assertEqual(row['Days Present'], 3)
        self.assertEqual(row['Days Late'], 1)
        self.assertEqual(row['Days Absent (Unexcused)'], 1)
        self.assertEqual(row['Days Absent (Excused)'], 1)
        self.assertEqual(row['Total Days Absent'], 2)
        self.assertEqual(row['Attendance %'], 60)

    def test_csv_zip(self):
        archive = zipfile.ZipFile(io.BytesIO(exports.csv_zip_bytes(YEAR, TERM, 5)))
        names = archive.namelist()
        for name in ('read_me.csv', 'learners.csv', 'parents.csv', 'marks_gr05_t3.csv',
                     'attendance_t3.csv', 'data_quality.csv'):
            self.assertIn(name, names)
        text = archive.read('learners.csv').decode('utf-8-sig')
        self.assertTrue(text.startswith('EMIS Number,Admission Number,LURITS Number'))
        self.assertIn('Mokoena', text)

    def test_management_command(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, 'export.xlsx')
            call_command('sasams_export', year=YEAR, term=TERM, grade=5, out=out, stdout=io.StringIO())
            self.assertIn('Learners', load_workbook(out).sheetnames)

    # -- pages --------------------------------------------------------------------
    def test_staff_only(self):
        urls = [reverse('sasams:index'), reverse('sasams:download'), reverse('sasams:learner-numbers')]
        for kind in ('student', 'educator'):
            self.client.force_login(person(f'{kind}-x@example.com', kind).user)
            for url in urls:
                self.assertEqual(self.client.get(url).status_code, 404, (kind, url))
        # (a parent without a linked child is sent to onboarding by middleware first;
        # either way they never see the export)
        self.client.force_login(person('parent-x@example.com', 'parent').user)
        for url in urls:
            self.assertIn(self.client.get(url).status_code, (302, 404), url)
        self.client.force_login(self.office.user)
        response = self.client.get(reverse('sasams:index'), {'year': YEAR, 'term': TERM, 'grade': 5})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Data quality')
        self.assertContains(response, 'fails the check-digit')
        response = self.client.get(reverse('sasams:download'), {'year': YEAR, 'term': TERM, 'grade': 5})
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertIn('Learners', load_workbook(io.BytesIO(response.content)).sheetnames)
        response = self.client.get(reverse('sasams:download'),
                                   {'year': YEAR, 'term': TERM, 'format': 'csv'})
        self.assertEqual(response['Content-Type'], 'application/zip')

    def test_learner_numbers_bulk_save(self):
        self.client.force_login(self.office.user)
        url = reverse('sasams:learner-numbers')
        page = self.client.get(url, {'year': YEAR, 'grade': 5})
        self.assertEqual(page.status_code, 200)
        self.assertContains(page, f'name="admission_{self.bad.pk}"')
        response = self.client.post(url, {'year': YEAR, 'grade': 5,
                                          f'admission_{self.bad.pk}': ' A0002 ',
                                          f'lurits_{self.bad.pk}': '987654321',
                                          f'admission_{self.good.pk}': 'A0001',
                                          f'lurits_{self.good.pk}': '123456789'})
        self.assertEqual(response.status_code, 302)
        self.bad.refresh_from_db()
        self.assertEqual((self.bad.admission_number, self.bad.lurits_number), ('A0002', '987654321'))

        # a number already used by another learner, and a non-numeric LURITS, are refused
        self.client.post(url, {'year': YEAR, 'grade': 5, f'admission_{self.bad.pk}': 'a0001',
                               f'lurits_{self.bad.pk}': 'ABC'})
        self.bad.refresh_from_db()
        self.assertEqual((self.bad.admission_number, self.bad.lurits_number), ('A0002', '987654321'))
