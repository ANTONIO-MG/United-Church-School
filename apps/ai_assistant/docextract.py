"""Plain-text extraction from uploaded .xlsx / .docx / .txt files.

Used by the AI authoring route to feed a memo/solution that isn't a PDF (e.g. a
solution workbook) to Claude as text context. PDFs are sent to Claude natively
as document blocks and do not go through here.
"""

import io
import os


def supported_memo(filename):
    """True if we can extract text from this file (or it's a PDF handled natively)."""
    ext = os.path.splitext(filename or '')[1].lower()
    return ext in ('.pdf', '.xlsx', '.docx', '.txt', '.md', '.csv')


def extract_text(uploaded):
    """Return plain text from an uploaded .xlsx/.docx/.txt/.csv/.md file (not PDF).

    Returns '' if the type is unsupported or extraction fails."""
    name = getattr(uploaded, 'name', '') or ''
    ext = os.path.splitext(name)[1].lower()
    try:
        raw = uploaded.read()
    except Exception:
        return ''
    if ext == '.xlsx':
        return _from_xlsx(raw)
    if ext == '.docx':
        return _from_docx(raw)
    if ext in ('.txt', '.md', '.csv'):
        try:
            return raw.decode('utf-8', errors='replace')
        except Exception:
            return ''
    return ''


def _from_xlsx(raw):
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        return ''
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    except Exception:
        return ''
    parts = []
    for ws in wb.worksheets:
        parts.append(f'### Sheet: {ws.title}')
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _from_docx(raw):
    try:
        import docx
    except ImportError:  # pragma: no cover
        return ''
    try:
        doc = docx.Document(io.BytesIO(raw))
    except Exception:
        return ''
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)
