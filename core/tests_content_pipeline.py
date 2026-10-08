"""Extraction, drafting and rubric-assisted marking — the document route.

The model is stubbed throughout. What is under test is everything *around* the
model, which is where the safety lives: that a pack is validated before it is
imported, that a bad draft is refused rather than half-landed, and that a
marking suggestion stays a suggestion.
"""

import io
import json
from decimal import Decimal
from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from apps.accounts.models import Person
from apps.assessments import models as assess_models
from apps.assessments.rubric_marking import SuggestionError, _shape, suggest, wants_suggestion
from core import content_pack
from core.content_pack import drafting
from core.content_pack.extract import ExtractionError, extract, extract_many
from core.testing import enrol, make_module, make_programme

User = get_user_model()


# ---------------------------------------------------------------------------
# Extraction — local, no model, no database
# ---------------------------------------------------------------------------
class ExtractionTests(TestCase):

    def test_plain_text_reads(self):
        text, meta = extract('Topic 5 — estate duty'.encode('utf-8'), 'notes.txt')
        self.assertIn('estate duty', text)
        self.assertEqual(meta['kind'], 'text')
        self.assertEqual(meta['filename'], 'notes.txt')

    def test_an_unknown_type_is_refused_by_name(self):
        with self.assertRaises(ExtractionError) as caught:
            extract(b'\x00\x01', 'scan.jpg')
        self.assertIn('.jpg', str(caught.exception))

    def test_an_empty_file_is_refused(self):
        with self.assertRaises(ExtractionError):
            extract(b'', 'guide.pdf')

    def test_a_file_of_whitespace_says_it_may_need_ocr(self):
        """A scanned PDF extracts to nothing, and the message has to say why —
        otherwise it looks like the importer is broken."""
        with self.assertRaises(ExtractionError) as caught:
            extract(b'   \n  ', 'scan.txt')
        self.assertIn('OCR', str(caught.exception))

    def test_an_oversized_file_is_refused_before_it_is_read(self):
        with self.assertRaises(ExtractionError) as caught:
            extract(b'x' * (41 * 1024 * 1024), 'huge.pdf')
        self.assertIn('limit', str(caught.exception))

    def test_a_docx_keeps_its_tables_as_rows(self):
        """The REQUIRED grid — the parts and their marks — is a table. Flattened
        into a paragraph it stops being readable as a grid."""
        try:
            import docx
        except ImportError:                                   # pragma: no cover
            self.skipTest('python-docx not installed')
        document = docx.Document()
        document.add_paragraph('Topic 5 mock')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '(a) Calculate the estate duty'
        table.cell(0, 1).text = '18'
        table.cell(1, 0).text = '(b) Explain s4(q)'
        table.cell(1, 1).text = '10'
        buffer = io.BytesIO()
        document.save(buffer)

        text, meta = extract(buffer.getvalue(), 'mock.docx')
        self.assertEqual(meta['kind'], 'docx')
        self.assertIn('(a) Calculate the estate duty | 18', text)
        self.assertIn('(b) Explain s4(q) | 10', text)

    def test_an_xlsx_keeps_its_mark_grid(self):
        try:
            import openpyxl
        except ImportError:                                   # pragma: no cover
            self.skipTest('openpyxl not installed')
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Solution'
        sheet.append(['REF', 'LINE', 'AUTHORITY', 'AMOUNT', 'MARKS'])
        sheet.append(['a', 'Primary residence', 's3(2)', 4500000, 1])
        buffer = io.BytesIO()
        book.save(buffer)

        text, meta = extract(buffer.getvalue(), 'solution.xlsx')
        self.assertEqual(meta['kind'], 'xlsx')
        self.assertIn('--- sheet: Solution ---', text)
        self.assertIn('Primary residence | s3(2) | 4500000 | 1', text)

    def test_one_bad_file_does_not_stop_the_others(self):
        text, metas, errors = extract_many([
            (b'good content here', 'guide.txt'),
            (b'', 'broken.pdf'),
            (b'more content', 'notes.md'),
        ])
        self.assertEqual(len(metas), 2)
        self.assertEqual(len(errors), 1)
        self.assertIn('good content here', text)
        self.assertIn('more content', text)


# ---------------------------------------------------------------------------
# Drafting — the model stubbed, everything around it real
# ---------------------------------------------------------------------------
GOOD_PACK = {
    'pack': 'thrive-pack/1',
    'module': 'IGNORED', 'programme': 'IGNORED',
    'topic': {'code': 'IGNORED', 'title': 'Whatever the model said'},
    'items': [{
        'type': 'mock', 'title': 'Topic 5 Mock', 'scenario_html': '<p>Johannes died.</p>',
        'total_marks': 5, 'minutes': 10,
        'parts': [{
            'ref': 'a', 'required': 'Calculate the estate duty', 'marks': 5,
            'marking': 'auto',
            'lines': [{'label': 'Primary residence', 'amount': 4500000, 'marks': 5}],
        }],
    }],
}


class DraftingTests(TestCase):

    def setUp(self):
        super().setUp()
        self.programme = make_programme('UCS', 'GR12')
        self.offering = make_module('Accounting', 'ACC', programme=self.programme)

    def draft(self, returned):
        with mock.patch.object(drafting, 'is_enabled', return_value=True), \
             mock.patch('apps.ai_assistant.generate.generate_from_prompt',
                        return_value=returned):
            return content_pack.draft_pack(
                'some extracted text', module='ACC', programme='UCS-GR12',
                topic_code='T5', topic_title='Estate duty')

    def test_the_placement_is_pinned_by_the_caller_not_the_model(self):
        """A wrong module code imports a TAX topic into auditing and looks like
        content until someone opens it. The caller decides, always."""
        pack, problems = self.draft(json.loads(json.dumps(GOOD_PACK)))
        self.assertEqual(pack['module'], 'ACC')
        self.assertEqual(pack['programme'], 'UCS-GR12')
        self.assertEqual(pack['topic']['code'], 'T5')
        self.assertEqual(pack['pack'], content_pack.PACK_VERSION)
        self.assertEqual(problems, [])

    def test_the_models_own_topic_title_is_kept(self):
        pack, _ = self.draft(json.loads(json.dumps(GOOD_PACK)))
        self.assertEqual(pack['topic']['title'], 'Whatever the model said')

    def test_a_draft_that_does_not_add_up_comes_back_with_the_problem(self):
        """The model misreading one row is the expected failure, and it must be
        reported rather than raised — a human fixes three figures, not the file."""
        bad = json.loads(json.dumps(GOOD_PACK))
        bad['items'][0]['parts'][0]['lines'][0]['marks'] = 3     # part says 5
        pack, problems = self.draft(bad)
        self.assertTrue(problems)
        self.assertTrue(any('allocate' in p for p in problems))
        self.assertIsNotNone(pack, 'the draft is still returned for correction')

    def test_a_non_object_response_is_an_error(self):
        with self.assertRaises(content_pack.DraftingError):
            self.draft(['not', 'an', 'object'])

    def test_drafting_is_refused_when_the_model_is_unavailable(self):
        with mock.patch.object(drafting, 'is_enabled', return_value=False):
            with self.assertRaises(content_pack.DraftingError) as caught:
                content_pack.draft_pack('text', module='ACC', programme='UCS-GR12',
                                        topic_code='T5', topic_title='Estate duty')
        self.assertIn('JSON route', str(caught.exception),
                      'the message must point at the route that still works')

    def test_a_drafted_pack_imports_through_the_ordinary_importer(self):
        pack, problems = self.draft(json.loads(json.dumps(GOOD_PACK)))
        self.assertEqual(problems, [])
        result = content_pack.import_pack(pack)
        self.assertEqual(result.topic.code, 'T5')
        paper = assess_models.Assessment.objects.get(source_ref__endswith='T5/mock')
        self.assertEqual(paper.status, assess_models.Assessment.STATUS_DRAFT,
                         'the document route publishes nothing')


# ---------------------------------------------------------------------------
# The import screen — administrators only, and JSON is the only door in
# ---------------------------------------------------------------------------
class ImportScreenTests(TestCase):

    def setUp(self):
        super().setUp()
        self.programme = make_programme('UCS', 'GR12')
        self.offering = make_module('Accounting', 'ACC', programme=self.programme)
        self.url = reverse('learning:content-import')

    def _user(self, email, role):
        user = User.objects.create_user(username=email, email=email, password='x')
        person = Person.objects.get(user=user)
        person.user_type = role
        person.registered = True
        person.profile_status = True
        person.save()
        return User.objects.get(pk=user.pk)

    def test_only_administrators_may_open_it(self):
        """403 for an educator or student; a parent is bounced earlier still by
        ParentAccessMiddleware, which redirects rather than forbids."""
        for role, expected in (('admin', 200), ('staff', 200),
                               ('educator', 403), ('student', 403)):
            user = self._user(f'{role}@example.com', role)
            self.client.force_login(user)
            self.assertEqual(self.client.get(self.url).status_code, expected, role)

        self.client.force_login(self._user('parent@example.com', 'parent'))
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def test_validating_reports_problems_and_imports_nothing(self):
        self.client.force_login(self._user('a@example.com', 'admin'))
        response = self.client.post(self.url, {
            'action': 'validate',
            'pack_json': json.dumps({'pack': 'wrong/9', 'items': []}),
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['problems'])
        self.assertFalse(assess_models.Assessment.objects.exists())

    def test_malformed_json_is_reported_not_raised(self):
        self.client.force_login(self._user('a@example.com', 'admin'))
        response = self.client.post(self.url, {'action': 'validate', 'pack_json': '{ nope'})
        self.assertEqual(response.status_code, 200)
        self.assertIn('Not valid JSON', response.context['problems'][0])

    def test_an_invalid_pack_cannot_be_imported(self):
        self.client.force_login(self._user('a@example.com', 'admin'))
        bad = json.loads(json.dumps(GOOD_PACK))
        bad['module'], bad['programme'] = 'ACC', 'UCS-GR12'
        bad['topic'] = {'code': 'T5', 'title': 'Estate duty'}
        bad['items'][0]['total_marks'] = 99          # parts add to 5
        response = self.client.post(self.url, {'action': 'import',
                                               'pack_json': json.dumps(bad)})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['problems'])
        self.assertFalse(assess_models.Assessment.objects.exists())

    def test_a_valid_pack_imports_as_draft(self):
        self.client.force_login(self._user('a@example.com', 'admin'))
        pack = json.loads(json.dumps(GOOD_PACK))
        pack.update(module='ACC', programme='UCS-GR12',
                    topic={'code': 'T5', 'title': 'Estate duty'})
        response = self.client.post(self.url, {'action': 'import',
                                               'pack_json': json.dumps(pack)})
        self.assertEqual(response.status_code, 302)
        paper = assess_models.Assessment.objects.get()
        self.assertEqual(paper.status, assess_models.Assessment.STATUS_DRAFT)


# ---------------------------------------------------------------------------
# Rubric assistance — a suggestion, never a mark
# ---------------------------------------------------------------------------
RUBRIC = [
    {'point': 's4(q) removes assets accruing to the spouse', 'authority': 's4(q)', 'marks': '3'},
    {'point': 'the abatement applies after s4 deductions', 'authority': 's4A', 'marks': '2'},
    {'point': 'the effect is a deferral, not an exemption', 'marks': '1'},
]


class RubricSuggestionTests(TestCase):

    def setUp(self):
        super().setUp()
        programme = make_programme('UCS', 'GR12')
        self.offering = make_module('Accounting', 'ACC', programme=programme)
        self.student = User.objects.create_user(username='s@x.com', email='s@x.com', password='x')
        person = Person.objects.get(user=self.student)
        person.user_type = 'student'
        person.save()
        enrol(person, self.offering)

        self.paper = assess_models.Assessment.objects.create(
            module=self.offering, title='Topic 5 Mock', kind='test', total_marks=6,
            status=assess_models.Assessment.STATUS_OPEN)
        section = assess_models.Section.objects.create(
            assessment=self.paper, title='Required', order=0)
        self.question = assess_models.Question.objects.create(
            section=section, type='long', order=0, text='(c) Explain s4(q).',
            marks=Decimal('6'),
            config={'kind': 'rubric', 'marking_intent': 'ai_assisted', 'rubric': RUBRIC},
            marking_mode=assess_models.Question.MARKING_MANUAL)
        attempt = assess_models.AssessmentAttempt.objects.create(
            assessment=self.paper, student=self.student,
            status=assess_models.AssessmentAttempt.STATUS_SUBMITTED)
        self.answer = assess_models.Answer.objects.create(
            attempt=attempt, question=self.question,
            response={'text': 'Section 4(q) takes out assets going to the spouse.'})

    def test_only_ai_assisted_parts_are_offered_assistance(self):
        self.assertTrue(wants_suggestion(self.question))

        self.question.config = dict(self.question.config, marking_intent='manual')
        self.assertFalse(wants_suggestion(self.question),
                         'professional-skills marks are a human judgement')

        self.question.config = {'kind': 'schedule', 'lines': []}
        self.assertFalse(wants_suggestion(self.question))

    def test_a_suggestion_adds_its_rows_up(self):
        result = _shape({'rows': [
            {'index': 0, 'verdict': 'made', 'quote': 'takes out assets going to the spouse'},
            {'index': 1, 'verdict': 'partial'},
            {'index': 2, 'verdict': 'missed'},
        ]}, RUBRIC)
        self.assertEqual(result['suggested'], Decimal('4.00'))   # 3 + 1 + 0
        self.assertEqual(result['total'], Decimal('6'))
        self.assertEqual([r['verdict'] for r in result['rows']],
                         ['made', 'partial', 'missed'])

    def test_a_row_the_model_left_out_counts_as_missed(self):
        """A suggestion that silently drops half the rubric would read as a low
        mark rather than as a failure."""
        result = _shape({'rows': [{'index': 0, 'verdict': 'made'}]}, RUBRIC)
        self.assertEqual(len(result['rows']), 3)
        self.assertEqual(result['rows'][2]['verdict'], 'missed')
        self.assertEqual(result['suggested'], Decimal('3.00'))

    def test_an_invented_row_index_is_dropped(self):
        result = _shape({'rows': [{'index': 99, 'verdict': 'made'}]}, RUBRIC)
        self.assertEqual(len(result['rows']), 3)
        self.assertEqual(result['suggested'], Decimal('0.00'))

    def test_an_unknown_verdict_is_treated_as_missed(self):
        result = _shape({'rows': [{'index': 0, 'verdict': 'brilliant'}]}, RUBRIC)
        self.assertEqual(result['rows'][0]['verdict'], 'missed')

    def test_an_empty_answer_is_refused_before_the_model_is_called(self):
        self.answer.response = {'text': ''}
        self.answer.save()
        with self.assertRaises(SuggestionError):
            suggest(self.answer)

    def test_a_question_with_no_rubric_is_refused(self):
        self.question.config = {}
        self.question.save(update_fields=['config'])
        with self.assertRaises(SuggestionError):
            suggest(self.answer)

    def test_a_suggestion_never_writes_a_mark(self):
        """The point of the whole feature: it advises, the coach decides."""
        stub = {'rows': [{'index': i, 'verdict': 'made'} for i in range(3)]}
        with mock.patch('apps.assessments.rubric_marking.is_enabled', return_value=True), \
             mock.patch('apps.ai_assistant.generate.generate_from_prompt', return_value=stub):
            result = suggest(self.answer)

        self.assertEqual(result['suggested'], Decimal('6.00'))
        self.answer.refresh_from_db()
        self.assertEqual(self.answer.awarded_marks, Decimal('0.00'))
        self.assertFalse(self.answer.marked)
